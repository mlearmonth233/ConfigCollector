"""Devices from an IT/OT Sitebook workbook.

The Sitebook's "Devices and Serials" tab lists every network device on a
site, one row per physical unit: column A the hostname, column C the
management IP, then PID, serial number, MAC, location and so on. A switch
stack is one row per member, every member with the stack's IP (the members
are named HOST, HOST-2, HOST-3...), so only the first row for each IP is a
device to manage; the rest are reported as stack members.
"""

from __future__ import annotations

import ipaddress
import re
from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import load_workbook

SHEET_NAME = "devices and serials"
# Column A (hostname) and C (IP) are the fixed layout; the rest are found
# by header name and are only there to help review.
NAME_COLUMN = 0
IP_COLUMN = 2
_OPTIONAL_HEADERS = {
    "location": ("site provided location", "location"),
    "pid": ("pid",),
    "serial": ("sn", "serial"),
}


class SitebookError(ValueError):
    pass


@dataclass
class SitebookDevice:
    row: int
    name: str
    host: str
    location: str | None = None
    pid: str | None = None
    serial: str | None = None
    stack_members: list[str] = field(default_factory=list)


@dataclass
class SitebookSkipped:
    row: int
    name: str | None
    host: str | None
    reason: str


@dataclass
class Sitebook:
    sheet: str
    devices: list[SitebookDevice]
    skipped: list[SitebookSkipped]


def _text(value) -> str:
    return "" if value is None else " ".join(str(value).split())


def _ip(value: str) -> str | None:
    candidate = value.strip()
    try:
        return str(ipaddress.ip_address(candidate))
    except ValueError:
        return None


def parse_sitebook(content: bytes) -> Sitebook:
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for a bad file
        raise SitebookError(f"Not an Excel workbook: {exc}") from exc
    try:
        ws = next((w for w in wb.worksheets if w.title.strip().lower() == SHEET_NAME), None)
        if ws is None:
            ws = next((w for w in wb.worksheets if "devices" in w.title.lower() and "serial" in w.title.lower()), None)
        if ws is None:
            raise SitebookError(f"No 'Devices and Serials' tab in this workbook (tabs: {', '.join(w.title for w in wb.worksheets)})")
        rows = list(ws.iter_rows(values_only=True, max_col=30))
    finally:
        wb.close()

    # The header row: "Host" in column A (usually row 1).
    header_index = next(
        (i for i, row in enumerate(rows[:10]) if row and _text(row[0]).lower() in ("host", "hostname", "host name")),
        None,
    )
    columns: dict[str, int] = {}
    if header_index is not None:
        labels = [_text(v).lower() for v in rows[header_index]]
        for key, names in _OPTIONAL_HEADERS.items():
            for wanted in names:
                hit = next((j for j, label in enumerate(labels) if label == wanted or label.startswith(wanted + " ")), None)
                if hit is not None:
                    columns[key] = hit
                    break
    start = 0 if header_index is None else header_index + 1

    devices: list[SitebookDevice] = []
    skipped: list[SitebookSkipped] = []
    by_ip: dict[str, SitebookDevice] = {}
    for number, row in enumerate(rows[start:], start=start + 1):
        cells = list(row) + [None] * 30
        name = _text(cells[NAME_COLUMN])
        raw_ip = _text(cells[IP_COLUMN])
        if not name and not raw_ip:
            continue
        if not name:
            skipped.append(SitebookSkipped(number, None, raw_ip or None, "no hostname in column A"))
            continue
        ip = _ip(raw_ip)
        if ip is None:
            skipped.append(SitebookSkipped(number, name, raw_ip or None, "no valid IP address in column C" if raw_ip else "no IP address in column C"))
            continue
        if ip in by_ip:
            first = by_ip[ip]
            first.stack_members.append(name)
            skipped.append(SitebookSkipped(number, name, ip, f"same IP as {first.name} (stack member)"))
            continue

        def col(key: str) -> str | None:
            return _text(cells[columns[key]]) or None if key in columns else None

        location = col("location")
        if location and location.lower() == "site to provide":
            location = _text(cells[1]) or None
        device = SitebookDevice(row=number, name=name, host=ip, location=location, pid=col("pid"), serial=col("serial"))
        by_ip[ip] = device
        devices.append(device)
    return Sitebook(sheet=ws.title, devices=devices, skipped=skipped)


# Model number -> (device_type, device_role), for rows the org's naming
# rules cannot place.
_PID_TYPES: tuple[tuple[re.Pattern, str, str | None], ...] = (
    (re.compile(r"^C9800|^CW9800", re.IGNORECASE), "cisco_wlc_9800", "wlc"),
    (re.compile(r"^AIR-CT|^AIR-WLC", re.IGNORECASE), "cisco_wlc", "wlc"),
    (re.compile(r"^FortiGate|^FG-|^FGT", re.IGNORECASE), "fortinet", "firewall"),
    (re.compile(r"^N\dK|^N9K|^Nexus", re.IGNORECASE), "cisco_nxos", "core_switch"),
    (re.compile(r"^AP[78]\d{3}", re.IGNORECASE), "apc_pdu", "pdu"),
    # Catalyst switches and Rockwell Stratix (1783-*, Cisco IOS / IOS-XE inside).
    (re.compile(r"^(?:WS-)?C\d{4}|^C9[2-6]\d{2}|^IE-?\d|^1783-", re.IGNORECASE), "cisco_ios", None),
    (re.compile(r"^ISR|^C8\d{3}|^ASR", re.IGNORECASE), "cisco_ios", "router"),
)


def type_from_pid(pid: str | None) -> tuple[str | None, str | None]:
    if not pid:
        return None, None
    for pattern, device_type, role in _PID_TYPES:
        if pattern.search(pid.strip()):
            return device_type, role
    return None, None
