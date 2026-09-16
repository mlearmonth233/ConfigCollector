"""The inventory as an Excel workbook: one sheet per table, a summary
sheet up front, headers frozen and filterable, columns sized to content."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.inventory import Inventory

_HEADER_FILL = PatternFill("solid", fgColor="0F766E")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _fmt_dt(value: datetime | None) -> str:
    return value.strftime("%Y-%m-%d %H:%M UTC") if value else ""


def _write_table(ws: Worksheet, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(["" if v is None else v for v in row])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for index, header in enumerate(headers, start=1):
        longest = max([len(str(header))] + [len(str(r[index - 1])) for r in rows if r[index - 1] is not None])
        ws.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 60)


def build_workbook(inventory: Inventory, org_name: str) -> bytes:
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    with_snapshot = [d for d in inventory.devices if d.has_snapshot]
    summary_rows = [
        ["Organization", org_name],
        ["Generated", _fmt_dt(inventory.generated_at)],
        ["Managed devices", len(inventory.devices)],
        ["  with a collected config", len(with_snapshot)],
        ["  with a serial number found", sum(1 for d in inventory.devices if d.serial)],
        ["Hardware components (show inventory)", len(inventory.hardware)],
        ["Access points", len(inventory.access_points)],
        ["Neighbour links (CDP/LLDP)", len(inventory.neighbors)],
        ["Unmanaged devices seen on the wire", len(inventory.unmanaged)],
        ["Endpoints (MAC addresses on access ports)", sum(1 for e in inventory.endpoints if not e.on_uplink)],
        [],
        ["Model", "Count"],
    ]
    summary_rows.extend([model, count] for model, count in inventory.models.items())
    for row in summary_rows:
        summary.append(row)
    summary["A1"].font = Font(bold=True)
    for cell in summary[12]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    summary.column_dimensions["A"].width = 46
    summary.column_dimensions["B"].width = 28

    _write_table(
        wb.create_sheet("Devices"),
        ["Name", "Hostname (reported)", "Management IP", "Type", "Site", "Role", "Model", "Serial", "Software", "Uptime", "Base MAC", "Config collected", "Missing commands"],
        [
            [d.name, d.hostname, d.host, d.device_type, d.site, d.role, d.model, d.serial, d.software, d.uptime, d.base_mac, _fmt_dt(d.collected_at) if d.has_snapshot else "never", ", ".join(d.commands_missing)]
            for d in inventory.devices
        ],
    )
    _write_table(
        wb.create_sheet("Hardware"),
        ["Device", "Component", "Description", "PID", "VID", "Serial"],
        [[h.device_name, h.name, h.description, h.pid, h.vid, h.serial] for h in inventory.hardware],
    )
    _write_table(
        wb.create_sheet("Neighbors"),
        ["Device", "Local port", "Neighbor", "Neighbor IP", "Platform", "Capabilities", "Remote port", "Protocol", "Managed in Packrat"],
        [[n.device_name, n.local_port, n.name, n.ip, n.platform, n.capabilities, n.remote_port, n.protocol.upper(), n.managed_device_name or "no"] for n in inventory.neighbors],
    )
    _write_table(
        wb.create_sheet("Access points"),
        ["Controller", "AP name", "Model", "MAC", "IP", "Serial", "Software"],
        [[a.controller_name, a.name, a.model, a.mac, a.ip, a.serial, a.software] for a in inventory.access_points],
    )
    _write_table(
        wb.create_sheet("Unmanaged"),
        ["Kind", "Name", "IP", "Platform", "Capabilities", "Seen from", "Protocol"],
        [[u.kind, u.name, u.ip, u.platform, u.capabilities, "; ".join(u.seen_from), "/".join(p.upper() for p in u.protocols)] for u in inventory.unmanaged],
    )
    _write_table(
        wb.create_sheet("Endpoints"),
        ["Switch", "Port", "VLAN", "MAC", "IP (from ARP)", "Entry type", "Port has a neighbour switch"],
        [[e.device_name, e.port, e.vlan, e.mac, e.ip, e.entry_type, "yes" if e.on_uplink else ""] for e in inventory.endpoints],
    )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
