"""Return to Operations (R2O): the hand-over checklist a controller must
pass before it goes back to operations, run against its latest snapshot.

The checklist is the organization's own workbook - one sheet per
controller build ("WLC AireOS", "Flex WLC 9800", "NEW WLC 9800"), one row
per check: a description, the command(s) to run, an expected output, the
pass/fail criteria in plain English, PRE or POST cutover, Mandatory or
Gold. It changes as the engineering standard does, so nothing about a
particular check is coded here. Instead each row's criteria are read for
the kinds of statement such checklists are written in, and every one
found becomes an expectation:

    Telnet must be "Disable"                 field = value
    v1 and v2c must be "Disable"             several fields, one value
    Channel Update Interval..... "12 Hours"  field = value (dotted, quoted)
    Product Version........ 8.10.162.0       field = value (dotted, first paragraph)
    "no ip http server" and "ip http secure-server" must be in output
    Words "public" and "private" not found in output
    ... must not contain "telnet" or "transport input all"
    Order must be "local" then "tacacs"
    Only those 2 destination IP should be configured   (from Expected value)
    Exporter IP must match depending on region: NA: a | EU: b   (any one)
    Must match the first 8 characters of device name
    Result must be as per expected value / Mentioned commands present
    Fail: No data returned

A row none of these fit is listed for manual review with the collected
output beside it. A few rows are about the access points themselves (AP
uplink speed, radios up, nothing in default-group); those reuse the
Wireless Audit's own AP checks. Every failure carries how to fix it: the
expectation that failed, any command the row itself names, and the lines
of the organization's configuration standard that set the same thing.

Commands are collected like any other: services/collector.py expands
"{site}" (the first 8 characters of the device name - the site code the
checklists write as XXXXXXX or <<SITECODE>>) and "{ap}" (once per AP).
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import load_workbook

from app.services import inventory as inv

SITE_PLACEHOLDER = "{site}"
AP_PLACEHOLDER = "{ap}"
MAX_OUTPUT_SNIPPET = 2500

_HEADERS = {
    "description": ("description",),
    "command": ("what and how", "command"),
    "expected": ("expected value",),
    "criteria": ("validation criteria", "pass/fail"),
    "automation": ("r2o automation",),
    "comments": ("comments",),
    "stage": ("pre or post",),
    "level": ("mandatory or gold",),
}


# --- importing the checklist -----------------------------------------------------------------


@dataclass
class R2OItem:
    item_id: str
    row: int
    description: str
    commands: list[str]
    expected: str
    criteria: str
    stage: str | None  # "PRE" | "POST"
    level: str | None  # "M" | "G"
    comments: str | None = None


@dataclass
class R2OSheet:
    name: str
    platform: str  # "9800" | "aireos"
    kind: str  # "flex" | "local"
    items: list[R2OItem]


def _text(value) -> str:
    return "" if value is None else str(value).replace("\r", "").strip()


_PROMPT = re.compile(r"^\s*\S+#\s*")


def split_commands(raw: str) -> list[str]:
    """The command cell -> the commands to run. Cells hold one command,
    several on separate lines, or a quoted list ("a" , "b"; "c"), sometimes
    with the device prompt pasted in front. Site-code and AP-name stand-ins
    become {site} / {ap}."""
    text = raw.replace("“", '"').replace("”", '"')
    parts: list[str] = []
    quoted = re.findall(r'"([^"]+)"', text)
    # A quoted list, unless the quotes are an argument (grep include Host "show logging").
    if quoted and not re.match(r"^\s*grep\b", text, re.IGNORECASE):
        rest = re.sub(r'"[^"]+"', "", text)
        parts = quoted + [p for p in re.split(r"[\n;,]+", rest) if p.strip()]
    else:
        parts = re.split(r"\n|;", text)
    commands = []
    for part in parts:
        cmd = _PROMPT.sub("", part.strip()).strip().strip(",").strip()
        if not cmd or cmd in ("-",) or not re.match(r"^(?:sh|show|grep|ping|dir|more|test)\b", cmd, re.IGNORECASE):
            continue
        cmd = re.sub(r"<<\s*site\s*code\s*>>|<<\s*site-?prefix\s*>>|(?<![A-Za-z0-9])X{5,}(?![A-Za-z0-9])", SITE_PLACEHOLDER, cmd, flags=re.IGNORECASE)
        cmd = re.sub(r"<\s*ap\s*name\s*>", AP_PLACEHOLDER, cmd, flags=re.IGNORECASE)
        cmd = " ".join(cmd.split())
        if cmd not in commands:
            commands.append(cmd)
    return commands


def _platform(name: str) -> str:
    return "aireos" if "aireos" in name.lower() else "9800"


def _kind(name: str) -> str:
    return "flex" if "flex" in name.lower() or "vwlc" in name.lower() else "local"


def import_checklist(content: bytes) -> list[R2OSheet]:
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheets: list[R2OSheet] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True, max_col=14))
        header_index = None
        columns: dict[str, int] = {}
        for i, row in enumerate(rows[:15]):
            labels = [_text(v).lower() for v in row]
            found = {key: j for key, needles in _HEADERS.items() for j, label in enumerate(labels) if label and any(n in label for n in needles)}
            if "command" in found and ("description" in found or "criteria" in found):
                header_index, columns = i, found
                break
        if header_index is None:
            continue
        items: list[R2OItem] = []
        for number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            row = list(row) + [None] * 14

            def col(key: str) -> str:
                return _text(row[columns[key]]) if key in columns else ""

            description = " ".join(col("description").split())
            commands = split_commands(col("command"))
            if not commands:
                continue
            stage = col("stage").upper()
            level = col("level").upper()
            items.append(
                R2OItem(
                    item_id=f"{ws.title}:{number}",
                    row=number,
                    description=description or commands[0],
                    commands=commands,
                    expected=col("expected"),
                    criteria=col("criteria"),
                    stage=stage if stage in ("PRE", "POST") else None,
                    level="M" if level.startswith("M") else ("G" if level.startswith("G") else None),
                    comments=col("comments") or None,
                )
            )
        if items:
            sheets.append(R2OSheet(name=ws.title, platform=_platform(ws.title), kind=_kind(ws.title), items=items))
    wb.close()
    return sheets


# --- expectations ------------------------------------------------------------------------------


@dataclass
class Expectation:
    kind: str  # field | present | absent | order | only | anyof | site_prefix | data | empty
    key: str | None = None
    value: str | None = None
    values: list[str] = field(default_factory=list)
    text: str = ""  # the criteria sentence it came from
    soft: bool = False  # from an example output / a per-region template: a warning, not a failure


_QUOTED = re.compile(r'"([^"]*)"')
_IP = r"\d{1,3}(?:\.\d{1,3}){3}"
# key ... must be / should match / needs to be / to be / also "value" (or "value2")
_STATEMENT = re.compile(
    r"(?P<key>[^\"]*?)\s*(?:\b(?:must|should|needs? to|has to)\s+(?:be\s+|match\s+|show\s+)?(?:set\s+to\s+)?|\bto\s+be\s+|\balso\s+)"
    r"\"(?P<value>[^\"]+)\"(?P<alts>(?:\s+or\s+\"[^\"]+\")*)",
    re.IGNORECASE,
)
_STATEMENT_BARE = re.compile(
    r"^(?P<key>[A-Za-z0-9][\w ./()-]{1,60}?)\s+(?:must|should)\s+be\s+(?P<value>[\w.+-]+(?:\s*(?:dBm|dB|Mbps|MHz|Hours?|Hour\(s\)|seconds|%))?)\s*\.?$",
    re.IGNORECASE,
)
_KEY = r"[A-Za-z0-9](?:[^.:=\n]|\.(?!\.))*?"
_DOTTED = re.compile(rf"^\s*(?P<key>{_KEY})\s*\.{{3,}}\s*(?P<value>\S.*?)\s*$")
_COLON = re.compile(rf"^\s*(?P<key>{_KEY})\s*(?::|\s=)\s+(?P<value>\S.*?)\s*$")
_FIELD_SEP = re.compile(rf"^\s*(?P<key>{_KEY})\s*(?:\.{{2,}}|:|=)\s*(?P<value>\S.*?)\s*$")
_VOLATILE_KEY = re.compile(r"time|uptime|count|number|last|since|mac|ip address|address|serial|age|seconds|ago|id\b|hash|version|date|clients|packets|bytes", re.IGNORECASE)
_VOLATILE_VALUE = re.compile(rf"^\d+(?:\.\d+)?$|{_IP}|(?:[0-9a-f]{{4}}\.){{2}}[0-9a-f]{{4}}|(?:[0-9a-f]{{2}}:){{5}}|\d+:\d+|ago|seconds|minutes|hours?,", re.IGNORECASE)


def _clean_key(key: str) -> str:
    key = re.sub(r"\([^)]*\)?", "", key)
    key = re.sub(r"^(?:the|and|also|then)\s+", "", key.strip(" ,.;"), flags=re.IGNORECASE)
    return " ".join(key.split()).strip(" .:-,")


def _keys(raw: str) -> list[str]:
    """'v1 and v2c' -> ['v1', 'v2c']; drops lead-ins ('For 5g-group01,')."""
    raw = re.sub(r"\([^)]*\)?", "", raw)
    raw = re.sub(r"^.*?\bfor\s+\S+\s*,\s*", "", raw, flags=re.IGNORECASE)
    keys = []
    for part in re.split(r"\s+and\s+|,\s*", raw):
        key = re.sub(r"\s+(?:status|value|output)$", "", _clean_key(part), flags=re.IGNORECASE)
        if key and len(key) <= 60 and not re.match(r"^(?:for|the remaining|remaining|below|it|they|this|both)\b", key, re.IGNORECASE):
            keys.append(key)
    return keys


_REGION_KEYS = {"na", "nala", "eu", "emea", "ap", "apac", "gc", "china", "la", "latam", "imea"}
_BARE_VALUE = re.compile(r"^(?:enabled?|disabled?|up|down|yes|no|none|on|off|[A-Z0-9_-]{2,}|-?\d+(?:\.\d+)?(?:\s*\S+)?)$", re.IGNORECASE)
_PROSE_WORDS = re.compile(r"\b(?:only|are|is|be|should|must|will|if|present|configured|output|verify|check)\b", re.IGNORECASE)


def _usable_field(key: str, value: str) -> bool:
    k = key.lower()
    if not key or not value or k in _REGION_KEYS or k in ("pass", "fail", "note", "e.g", "example"):
        return False
    if re.search(r"\b(?:output|should|must|find|see)\b", k) or value.lower().startswith(("show ", "sh ")):
        return False
    return not ("|" in value and re.search(r"\b(?:EU|AP|NA|GC)\s*:", value))


def _sample_site_codes(item: R2OItem) -> set[str]:
    """The site code the checklist's example happens to use (SITE0001 in
    'show wireless tag rf detailed SITE0001-APGROUP01'), so it can be
    swapped for the audited controller's own."""
    codes: set[str] = set()
    text = item.expected + "\n" + item.criteria
    for command in item.commands:
        if SITE_PLACEHOLDER not in command:
            continue
        pattern = re.escape(command).replace(re.escape(SITE_PLACEHOLDER), r"([A-Z0-9]{8})")
        pattern = pattern.replace(r"\ ", r"\s+")
        for m in re.finditer(pattern, text, re.IGNORECASE):
            codes.add(m.group(1))
    if re.search(r"first 8 characters of (?:the )?device name", item.criteria, re.IGNORECASE):
        codes.update(re.findall(r"\b([A-Z]{4}[A-Z0-9]{4})(?=\b|_)", item.criteria + "\n" + item.expected))
    return codes


def _siteify(value: str, codes: set[str]) -> str:
    for code in codes:
        value = re.sub(re.escape(code), SITE_PLACEHOLDER, value)
    return value


def expectations(item: R2OItem) -> list[Expectation]:
    crit = item.criteria.replace("“", '"').replace("”", '"').replace("’", "'")
    low = crit.lower()
    found: list[Expectation] = []
    codes = _sample_site_codes(item)
    regional = bool(re.search(r"per region|fits the template|vary per region", low + " " + item.description.lower()))

    if re.search(r"first 8 characters of (?:the )?device name", low):
        found.append(Expectation("site_prefix", text="Must match the first 8 characters of the device name"))
    m = re.search(r"only (?:those|these)\s+(?:\d+\s+)?(?P<what>[\w ]+?)(?:\s*\((?P<names>[^)]*)\))?\s+(?:should|must|are)", crit, re.IGNORECASE)
    if m:
        names = [n.strip() for n in re.split(r",|\band\b", m.group("names") or "") if n.strip()]
        if names:
            found.append(Expectation("only", key="name", values=names, text=m.group(0)))
        elif re.search(r"\bip\b|destination|host|server", m.group("what"), re.IGNORECASE):
            ips = re.findall(_IP, item.expected)
            if ips:
                found.append(Expectation("only", key="ip", values=list(dict.fromkeys(ips)), text=m.group(0)))
    m = re.search(r"only\s+(?P<names>[\w-]+(?:\s*(?:,|and)\s*[\w-]+)+)\s+(?:are|is)\s+present", crit, re.IGNORECASE)
    if m:
        found.append(Expectation("only", key="name", values=[n.strip() for n in re.split(r",|\band\b", m.group("names")) if n.strip()], text=m.group(0)))
    if re.search(r"depending on region|\(1 of them\)|region:|for na\b.*for eu", low):
        ips = list(dict.fromkeys(re.findall(_IP, crit)))
        if ips and not any(e.kind == "only" for e in found):
            found.append(Expectation("anyof", values=ips, text="One of the regional values: " + ", ".join(ips)))
    m = re.search(r"order must be\s+((?:\"[^\"]+\"\s*(?:then|,|and)?\s*)+)", crit, re.IGNORECASE)
    if m:
        found.append(Expectation("order", values=_QUOTED.findall(m.group(1)), text=m.group(0).strip()))
    if re.search(r"no data returned", low):
        found.append(Expectation("data", text="The command must return data"))
    if re.search(r"pass:?\s*(?:if there is )?\(?no output|pass if there is no output", low):
        found.append(Expectation("empty", text="PASS if there is no output"))

    mode = None  # None | "require" | "forbid"
    first_paragraph = True
    for raw in crit.splitlines():
        line = " ".join(raw.split())
        if not line:
            if found or mode:
                first_paragraph = False
            continue
        ll = line.lower()
        if _PROMPT.match(raw) and "#" in line.split()[0]:
            continue  # pasted prompt / command line
        if re.match(r"^(?:pass|fail)\b", ll):
            if ll.startswith("fail") and re.search(r"output of|if there is|is enabled|present", ll):
                mode = "forbid"
            elif ll.startswith("pass") and re.search(r"contains|must match|should match", ll):
                mode = "require"
            else:
                mode = None
            rest = re.sub(r"^(?:pass|fail)\s*:?\s*", "", line, flags=re.IGNORECASE)
            if not rest or rest.endswith(":") or re.search(r"no output|contains|output of", rest, re.IGNORECASE):
                continue
            line, ll = rest, rest.lower()
        if re.search(r"(?:should|must) match\s*:?\s*$|requirements:?\s*$|below .*must match|contains the parameters|policy contains:?$", ll):
            mode = "require"
            continue
        if mode == "forbid":
            if not re.search(r"\b(?:if|else)\b", ll):
                found.append(Expectation("absent", value=line, text="FAIL if present: " + line))
            continue
        if mode == "require":
            fm = _DOTTED.match(raw) or _COLON.match(raw)
            if fm and fm.group("value").strip():
                found.append(Expectation("field", key=_clean_key(fm.group("key")), value=_siteify(fm.group("value").strip().strip('"'), codes), text=line))
            elif not line.endswith(":") and not re.match(r"^-{3,}", line):
                found.append(Expectation("present", value=_siteify(line, codes), text=line))
            continue

        # Sentence by sentence.
        for sentence in re.split(r"(?<=[.!])\s+(?=[A-Z\"])", line):
            sl = sentence.lower()
            quoted = _QUOTED.findall(sentence)
            cm = re.search(r"(?P<name>[\w-]+) must be present and must be \"(?P<value>[^\"]+)\" under \"(?P<col>[^\"]+)\" column", sentence, re.IGNORECASE)
            if cm:
                found.append(Expectation("present", value=f"{cm.group('name')} {cm.group('value')}", text=sentence))
                continue
            if quoted and re.search(r"not found|must not|should not|cannot|can't", sl):
                negative_at = re.search(r"not found|must not|should not|cannot|can't", sl).start()
                for qm in _QUOTED.finditer(sentence):
                    if "not found" in sl or qm.start() > negative_at:
                        found.append(Expectation("absent", value=qm.group(1), text=sentence))
                    else:
                        found.append(Expectation("present", value=qm.group(1), text=sentence))
                continue
            if quoted and re.search(r"must be (?:present|visible|in (?:the )?output)|is present in output|should be present|must exist|must be in output|showing there", sl):
                found += [Expectation("present", value=_siteify(q, codes), text=sentence) for q in quoted if q.strip()]
                continue
            if sl.startswith("order must be"):
                continue
            statements = list(_STATEMENT.finditer(sentence))
            if statements:
                pos_key = None
                for st in statements:
                    keys = _keys(st.group("key")) or ([pos_key] if pos_key else [])
                    values = [st.group("value")] + _QUOTED.findall(st.group("alts") or "")
                    both = re.search(r"for both (\w+) and (\w+)\s*(\w+)?", sentence[st.end():], re.IGNORECASE)
                    if both:
                        noun = f" {both.group(3)}" if both.group(3) else ""
                        keys = [f"{both.group(1)}{noun}", f"{both.group(2)}{noun}"]
                    for key in keys:
                        if len(values) > 1:
                            found.append(Expectation("field", key=key, values=[_siteify(v.strip(), codes) for v in values], value=values[0].strip(), text=sentence))
                        else:
                            found.append(Expectation("field", key=key, value=_siteify(values[0].strip(), codes), text=sentence))
                    pos_key = keys[-1] if keys else pos_key
                continue
            if quoted and re.search(r"\b(?:should|must)\b", sl) and len(quoted) == 1 and quoted[0].strip():
                found.append(Expectation("present", value=_siteify(quoted[0], codes), text=sentence))
                continue
            fm = _DOTTED.match(raw) or (_COLON.match(raw) if first_paragraph else None)
            if fm and fm.group("value").strip():
                value = fm.group("value").strip()
                if re.search(r"\b(?:should|must)\b", fm.group("key"), re.IGNORECASE):
                    key_text = re.split(r"\b(?:should|must)\b", fm.group("key"), flags=re.IGNORECASE)[0]
                    found.append(Expectation("field", key=_clean_key(key_text), value=_siteify(value.strip('"'), codes), text=line))
                    break
                q = _QUOTED.fullmatch(value)
                if q or first_paragraph:
                    found.append(Expectation("field", key=_clean_key(fm.group("key")), value=_siteify((q.group(1) if q else value).strip('"'), codes), text=line))
                break
            bm = _STATEMENT_BARE.match(sentence)
            if bm and not quoted and not sl.startswith("only") and _BARE_VALUE.match(bm.group("value")):
                found.append(Expectation("field", key=_clean_key(bm.group("key")), value=bm.group("value"), text=sentence))
                continue
            # A bare configuration line as the whole criterion (" virtual-ip ipv4 192.0.2.1").
            if first_paragraph and len(crit.strip().splitlines()) <= 3 and raw.startswith(" ") is not None and re.match(r"^[a-z0-9]", line) and not _PROSE_WORDS.search(sl):
                found.append(Expectation("present", value=_siteify(line, codes), text=line))

    # "As per expected value" style, or an Expected value with no criteria:
    # the Expected value's lines (configuration) or fields (show output).
    as_expected = re.search(r"as per expected value|mentioned commands present|like presented here|must match the text shown|fits the template|commands present in output", low)
    if as_expected or (not crit.strip() and item.expected.strip()):
        config_like = bool(as_expected)
        soft = regional or not as_expected
        for raw in item.expected.splitlines():
            line = " ".join(raw.split()).strip('"')
            if not line or line in ("!", "-") or len(line) < 4 or ("#" in line.split()[0] and _PROMPT.match(raw)):
                continue
            if re.match(r"^(?:sh|show)\s", line, re.IGNORECASE) or line.endswith(":") or _QUOTED.search(line) and _PROSE_WORDS.search(line):
                continue
            if config_like:
                found.append(Expectation("present", value=_siteify(line, codes), text="Expected value: " + line, soft=soft))
                continue
            fm = _COLON.match(raw) or _DOTTED.match(raw)
            if fm:
                key, value = _clean_key(fm.group("key")), fm.group("value").strip()
                if value and len(value) <= 40 and not _VOLATILE_KEY.search(key) and not _VOLATILE_VALUE.search(value):
                    found.append(Expectation("field", key=key, value=_siteify(value, codes), text="Example output: " + line, soft=True))
    if regional:
        for e in found:
            if e.kind == "present":
                e.soft = True

    unique: dict[tuple, Expectation] = {}
    for e in found:
        if e.kind == "field" and not (e.key and (e.value or e.values) and _usable_field(e.key, e.value or "")):
            continue
        unique.setdefault((e.kind, (e.key or "").lower(), (e.value or "").lower(), tuple(e.values)), e)
    return list(unique.values())


# --- evaluating --------------------------------------------------------------------------------


def _norm(text: str) -> str:
    return " ".join(text.lower().split())


_SYNONYMS = {"disable": "disabled", "enable": "enabled", "readonly": "read only", "readwrite": "read write", "yes": "yes", "up": "up"}


def _same(expected: str, actual: str) -> bool:
    e, a = _norm(expected.replace("/", " ")).strip('"'), _norm(actual.replace("/", " "))
    if not e:
        return True
    if a == e or a.startswith(e) or f" {e}" in f" {a} ":
        return True
    e2, a2 = _SYNONYMS.get(e, e), _SYNONYMS.get(a.split()[0], a.split()[0]) if a else ""
    return e2 == a2 or e2.rstrip("d") == a2.rstrip("d")


def find_field(output: str, key: str) -> str | None:
    """Value of `key` in "Key ....... value" / "Key : value" / "Key = value"
    output. An exact key wins (case-sensitive first: AireOS prints both
    "802.1X" and "802.1x"), then a key the wanted one is a whole-word part
    of ("v1" in "SNMP v1 Mode")."""
    want_raw = " ".join(re.sub(r"\(.*?\)", "", key).split())
    want = want_raw.lower()
    if not want:
        return None
    rows = []
    for raw in output.splitlines():
        m = _FIELD_SEP.match(raw)
        if m:
            k = " ".join(re.sub(r"\(.*?\)", "", m.group("key")).split())
            rows.append((k, m.group("value").strip()))
    for k, v in rows:
        if k == want_raw:
            return v
    for k, v in rows:
        if k.lower() == want:
            return v
    for k, v in rows:
        kl = k.lower()
        if kl.startswith(want + " ") or kl.endswith(" " + want) or f" {want} " in f" {kl} ":
            return v
    return None


def _contains(output: str, needle: str) -> bool:
    n = _norm(needle)
    if "<" in n and ">" in n:  # "key <secretkey>": a placeholder value
        pattern = r"\s+".join(r"\S+" if re.fullmatch(r"<[^>]+>", w) else re.escape(w) for w in n.split())
        return re.search(pattern, _norm(output)) is not None
    return n in _norm(output)


_ERROR_OUTPUT = re.compile(r"^\s*(?:%\s*Invalid|Incorrect usage|% Incomplete|% Unknown|Invalid input)", re.IGNORECASE | re.MULTILINE)


@dataclass
class R2OResult:
    item: R2OItem
    status: str  # pass | fail | warn | manual | no_data | na
    details: list[str]
    output: str
    commands_run: list[str]
    fix: list[str]


def expand(command: str, device_name: str, ap_names: list[str] | None = None) -> list[str]:
    site = site_code(device_name)
    cmd = command.replace(SITE_PLACEHOLDER, site)
    if AP_PLACEHOLDER in cmd:
        return [cmd.replace(AP_PLACEHOLDER, ap) for ap in (ap_names or [])]
    return [cmd]


def site_code(device_name: str) -> str:
    return device_name.strip()[:8]


def _output_for(content: str, command: str) -> str | None:
    for typed, output in inv.sections(content):
        if typed and _norm(typed) == _norm(command):
            return output
    # Fall back to the abbreviation-tolerant match (no pipes in the command).
    if "|" not in command:
        return inv.find_section(content, command)
    return None


def evaluate(item: R2OItem, content: str, device_name: str, ap_names: list[str], ap_findings: dict[str, list[dict]], fix_lines) -> R2OResult:
    """One checklist row against a controller's snapshot. `ap_findings`
    are the Wireless Audit's findings by check id (for the AP rows);
    `fix_lines(item)` returns lines of the configuration standard that set
    the same thing."""
    run: list[str] = []
    outputs: list[str] = []
    missing: list[str] = []
    for command in item.commands:
        for cmd in expand(command, device_name, ap_names):
            run.append(cmd)
            out = _output_for(content, cmd)
            if out is None:
                missing.append(cmd)
            else:
                outputs.append(out)
    output = "\n".join(outputs)
    fix = fix_lines(item)
    severity = "warn" if item.level == "G" else "fail"

    special = _special(item, ap_findings, output, severity)
    if special is not None:
        status, details = special
        return R2OResult(item, status, details, output[:MAX_OUTPUT_SNIPPET], run, fix if status in ("fail", "warn") else [])

    if not outputs:
        return R2OResult(item, "no_data", [f"Not collected: {', '.join(missing) or ', '.join(item.commands)}. Collect the controllers from the Wireless Audit page (it adds the R2O commands), then run the audit again."], "", run, [])

    exps = expectations(item)
    if not exps:
        return R2OResult(item, "manual", ["No machine-checkable criteria in this row: compare the output with the expected value by eye.", *([item.criteria[:400]] if item.criteria else [])], output[:MAX_OUTPUT_SNIPPET], run, [])

    failures: list[str] = []
    soft_failures: list[str] = []
    if _ERROR_OUTPUT.search(output) and len(output.strip().splitlines()) <= 3:
        failures.append(f"The controller rejected the command: {output.strip()[:200]}")
    for e in exps:
        problem = _check(e, output, device_name, item)
        if problem:
            (soft_failures if e.soft else failures).append(problem)
    if missing and not failures and not soft_failures:
        return R2OResult(item, "no_data", [f"Some commands were not collected: {', '.join(missing)}"], output[:MAX_OUTPUT_SNIPPET], run, [])
    if failures:
        status = severity
    elif soft_failures:
        status = "warn"
        soft_failures.insert(0, "Differs from the checklist's example / regional template (review):")
    else:
        status = "pass"
    details = failures + soft_failures
    return R2OResult(item, status, details or [f"{len(exps)} criteria met"], output[:MAX_OUTPUT_SNIPPET], run, fix if details else [])


def _check(e: Expectation, output: str, device_name: str, item: R2OItem) -> str | None:
    if e.kind == "field":
        wanted = [v.replace(SITE_PLACEHOLDER, site_code(device_name)) for v in (e.values or [e.value or ""])]
        shown = " or ".join(f"'{w}'" for w in wanted)
        actual = find_field(output, e.key or "")
        if actual is None:
            # A table cell ("1  0  192.168.16.1  In Sync") or "key value" row.
            if any(_contains(output, f"{e.key} {w}") or re.search(rf"(?i)(?<![\w-]){re.escape(w)}(?![\w-])", output) for w in wanted):
                return None
            return f"{e.key}: not found in the output (expected {shown})"
        return None if any(_same(w, actual) for w in wanted) else f"{e.key} is '{actual}', expected {shown}"
    if e.kind == "empty":
        lines = [line for line in output.splitlines() if line.strip() and not _PROMPT.match(line + " ")]
        return None if not lines else "Expected no output, found: " + "; ".join(line.strip() for line in lines[:5])
    if e.kind == "present":
        value = (e.value or "").replace(SITE_PLACEHOLDER, site_code(device_name))
        return None if _contains(output, value) else f"Missing from the output: {value}"
    if e.kind == "absent":
        return f"Must not be present, but found: {e.value}" if re.search(rf"(?i)\b{re.escape(e.value or '')}\b", output) else None
    if e.kind == "order":
        positions = [(_norm(output).find(_norm(v)), v) for v in e.values]
        if any(p < 0 for p, _ in positions):
            return f"Not all of {', '.join(e.values)} are in the output"
        return None if [v for _, v in sorted(positions)] == e.values else f"Order is {' then '.join(v for _, v in sorted(positions))}, expected {' then '.join(e.values)}"
    if e.kind == "only":
        if e.key == "ip":
            seen = set(re.findall(_IP, output))
            extra = sorted(seen - set(e.values))
            gone = [v for v in e.values if v not in seen]
        else:
            names = _first_column_names(output)
            wanted = {v.lower() for v in e.values}
            extra = sorted(n for n in names if n.lower() not in wanted)
            gone = [v for v in e.values if v.lower() not in {n.lower() for n in names}]
        parts = []
        if extra:
            parts.append("not in the standard: " + ", ".join(extra))
        if gone:
            parts.append("missing: " + ", ".join(gone))
        return f"Only {', '.join(e.values)} expected - " + "; ".join(parts) if parts else None
    if e.kind == "anyof":
        return None if any(v in output for v in e.values) else f"None of the expected values found ({', '.join(e.values)})"
    if e.kind == "site_prefix":
        code = site_code(device_name)
        return None if re.search(rf"(?i)\b{re.escape(code)}", output) else f"Expected the site code {code} (first 8 characters of {device_name}) in the output"
    if e.kind == "data":
        return None if output.strip() and not _ERROR_OUTPUT.search(output) else "No data returned"
    return None


def _first_column_names(output: str) -> list[str]:
    """User/account names from a table (AireOS show mgmtuser) or IOS-XE
    'username X ...' lines."""
    names = re.findall(r"(?im)^\s*username\s+(\S+)", output)
    if names:
        return list(dict.fromkeys(names))
    rows = []
    started = False
    for line in output.splitlines():
        if re.match(r"^\s*-{3,}", line):
            started = True
            continue
        if started and line.strip():
            rows.append(line.split()[0])
    return list(dict.fromkeys(rows))


def _special(item: R2OItem, ap_findings: dict[str, list[dict]], output: str, severity: str) -> tuple[str, list[str]] | None:
    """Rows about the access points reuse the audit's own AP checks."""
    cmds = " ".join(item.commands).lower()
    desc = item.description.lower()
    notes = f"{item.criteria} {item.comments or ''}".lower()

    if "redundancy" in cmds and output and re.search(r"2 wlc|only 1 wlc|only one|skip this test", notes):
        if re.search(r"SSO DISABLED|Hardware Mode\s*=\s*Simplex|Peer State\s*=\s*N/A|Redundancy Mode \(Operational\)\s*=\s*Non-redundant", output, re.IGNORECASE):
            return "na", ["Standalone controller (no SSO peer): the checklist skips this check."]

    def from_findings(check_ids: tuple[str, ...], ok: str, filt=None) -> tuple[str, list[str]]:
        found = [f for cid in check_ids for f in ap_findings.get(cid, []) if filt is None or filt(f)]
        if not found:
            return "pass", [ok]
        return severity, [f"{f['subject']}: {f['message']}" for f in found[:50]] + ([f"... and {len(found) - 50} more"] if len(found) > 50 else [])

    if "ethernet" in cmds and ("speed" in desc or "speed" in item.criteria.lower() or "connected with" in desc):
        return from_findings(("ap.uplink",), "Every AP uplink is up at the expected speed")
    if re.search(r"802\.11-abgn summary|802\.11b summary|24ghz summary", cmds) and "operation" in desc:
        return from_findings(("radio.state",), "Every 2.4 GHz radio is up", lambda f: "2.4 GHz" in f["message"])
    if re.search(r"802\.11a summary|5ghz summary", cmds) and "operation" in desc:
        return from_findings(("radio.state",), "Every 5 GHz radio is up", lambda f: "5 GHz" in f["message"])
    if "apgroups" in cmds and "default-group" in item.criteria.lower():
        return from_findings(("grp.default",), "No AP in default-group")
    if "apgroups" in cmds and "format" in item.criteria.lower():
        return from_findings(("grp.naming",), "Every AP group follows the naming format")
    if "ap tag summary" in cmds and "default" in item.criteria.lower():
        return from_findings(("grp.default", "ap.tags", "grp.mixed"), "No AP on a default tag, no misconfigured tags, generations kept apart")
    if re.match(r"(?:show|sh) ap summary", cmds.strip()) and ("quantity" in desc or "count" in desc):
        count = re.search(r"Number of APs\s*[.:]*\s*(\d+)", output)
        return "manual", [f"{count.group(1) if count else 'unknown'} APs registered: compare with the CMDB / site documentation."]
    if "inventory" in cmds and "serial" in desc:
        serials = re.findall(r"SN:\s*(\S+)", output)
        return "manual", [f"Serial number(s) {', '.join(serials) or 'not found'}: compare with the CMDB."]
    return None


# --- how to fix ----------------------------------------------------------------------------------

_STOP = {"show", "sh", "run", "running-config", "include", "inc", "in", "i", "s", "section", "b", "begin", "detailed", "detail", "summary", "all", "name", "|", "grep", "{site}", "{ap}", "and", "the"}


def standard_lines(item: R2OItem, template: str, platform: str, limit: int = 10) -> list[str]:
    """Lines of the configuration standard that configure what a row
    checks: AireOS "show advanced fra sensor" -> "config advanced fra ...";
    IOS-XE "show run | i ntp" -> the template's ntp lines."""
    if not template:
        return []
    lines = [line.rstrip() for line in template.splitlines() if line.strip() and not line.lstrip().startswith("#")]
    out: list[str] = []
    for command in item.commands:
        cmd = command.lower()
        if platform == "aireos":
            words = [w for w in re.sub(r"^(?:show|sh)\s+", "", cmd).split() if w not in _STOP]
            for n in (min(3, len(words)), 2, 1):
                if n <= 0:
                    continue
                prefix = "config " + " ".join(words[:n])
                hits = [line for line in lines if re.sub(r"^(?:optional|absent):\s*", "", line.strip().lower()).startswith(prefix)]
                if hits:
                    out += hits
                    break
        else:
            pipe = re.search(r"\|\s*(?:i|in|inc|include|s|section|b|begin)\s+(.+)$", cmd)
            if pipe:
                needle = pipe.group(1).strip().strip('"')
                out += [line for line in lines if needle in line.lower()]
            else:
                words = [w for w in cmd.split() if w not in _STOP and not w.startswith("{")]
                for word in words[:3]:
                    out += [line for line in lines if not line.startswith(" ") and re.search(rf"(?i)^(?:optional: |absent: )?(?:\S+ )?{re.escape(word)}\b", line.strip())]
        if len(out) >= limit:
            break
    seen: list[str] = []
    for line in out:
        if line not in seen:
            seen.append(line)
    return seen[:limit]


_FIX_HINTS: list[tuple[re.Pattern, dict[str, str]]] = [
    (re.compile(r"telnet", re.I), {"aireos": "config network telnet disable", "9800": "line vty 0 15 / transport input ssh"}),
    (re.compile(r"\bhttp\b|web mode", re.I), {"aireos": "config network webmode disable ; config network secureweb enable", "9800": "no ip http server ; ip http secure-server"}),
    (re.compile(r"snmp v1|snmp v2|v1 and v2", re.I), {"aireos": "config snmp version v1 disable ; config snmp version v2c disable ; config snmp version v3 enable", "9800": "no snmp-server community <name> (use SNMPv3 users only)"}),
    (re.compile(r"communit", re.I), {"aireos": "config snmp community delete public ; config snmp community delete private", "9800": "no snmp-server community public ; no snmp-server community private"}),
    (re.compile(r"\bfra\b|flexible radio", re.I), {"aireos": "config advanced fra disable", "9800": "ap fra (disable in the RF configuration: no ap fra)"}),
    (re.compile(r"clock|ntp", re.I), {"aireos": "config time ntp server <index> <ip> ; check reachability of the NTP server", "9800": "ntp server <ip> ; ntp source <mgmt interface>"}),
    (re.compile(r"priority order|aaa auth", re.I), {"aireos": "config aaa auth mgmt local tacacs"}),
    (re.compile(r"vlan ?1", re.I), {"9800": "interface Vlan1 / shutdown"}),
    (re.compile(r"ssh v2|ssh version", re.I), {"9800": "ip ssh version 2"}),
    (re.compile(r"licens", re.I), {"aireos": "license smart register idtoken <token>", "9800": "license smart trust idtoken <token> all ; check 'show license status'"}),
    (re.compile(r"optimi[sz]ed roam", re.I), {"9800": "no ap dot11 24ghz optimized-roaming ; no ap dot11 5ghz optimized-roaming"}),
    (re.compile(r"mgmt-via-wireless|wireless mgmt", re.I), {"9800": "no wireless mgmt-via-wireless"}),
]


def fix_hints(item: R2OItem, platform: str) -> list[str]:
    hints = []
    # A command the row names itself: "(config aaa auth mgmt local tacacs)".
    for m in re.finditer(r"\((config [^)]+)\)", item.description + " " + item.criteria):
        hints.append(m.group(1).strip())
    for pattern, by_platform in _FIX_HINTS:
        if pattern.search(item.description) and platform in by_platform:
            hints.append(by_platform[platform])
    return list(dict.fromkeys(hints))
