"""Inventory: parsing hardware facts out of collected snapshots, joining
them into one inventory, the API, and the Excel export."""

from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app import tasks as tasks_module
from app.services import inventory as inv

pytestmark = pytest.mark.asyncio


# --- sample outputs (abridged real-world layouts) -------------------------------------------

IOS_SHOW_VERSION = """Cisco IOS XE Software, Version 17.09.04a
Cisco IOS Software [Cupertino], Catalyst L3 Switch Software (CAT9K_IOSXE), Version 17.9.4a, RELEASE SOFTWARE (fc3)
Technical Support: http://www.cisco.com/techsupport

HQ-CORE-SW01 uptime is 1 year, 12 weeks, 3 days, 4 hours, 12 minutes
Uptime for this control processor is 1 year, 12 weeks, 3 days, 4 hours, 14 minutes
System returned to ROM by Reload Command

cisco C9300-48P (X86) processor with 1338114K/6147K bytes of memory.
Processor board ID FOC2345X0AB
Base Ethernet MAC Address          : 70:0f:6a:11:22:00
Model Number                       : C9300-48P
System Serial Number               : FOC2345X0AB
"""

NXOS_SHOW_VERSION = """Cisco Nexus Operating System (NX-OS) Software
Software
  BIOS: version 05.47
  NXOS: version 9.3(12)
Hardware
  cisco Nexus9000 C93180YC-EX Chassis
  Intel(R) Xeon(R) CPU  @ 1.80GHz with 24571952 kB of memory.
  Processor Board ID FDO23456ABC

  Device name: DC1-LEAF-01
  bootflash:   53298520 kB
Kernel uptime is 200 day(s), 3 hour(s), 21 minute(s), 9 second(s)
"""

SHOW_INVENTORY = """NAME: "c93xx Stack", DESCR: "c93xx Stack"
PID: C9300-48P         , VID: V02  , SN: FOC2345X0AB

NAME: "Switch 1", DESCR: "C9300-48P"
PID: C9300-48P         , VID: V02  , SN: FOC2345X0AB

NAME: "Switch 1 - Power Supply A", DESCR: "Switch 1 - Power Supply A"
PID: PWR-C1-715WAC-P   , VID: V01  , SN: DTN2340V0AB

NAME: "TenGigabitEthernet1/1/1", DESCR: "SFP-10GBase-SR"
PID: SFP-10G-SR          , VID: V03  , SN: AVD2345A0AB
"""

CDP_DETAIL = """-------------------------
Device ID: HQ-DIST-SW11.corp.local
Entry address(es):
  IP address: 10.10.1.11
Platform: cisco C9200L-48P-4X,  Capabilities: Switch IGMP
Interface: GigabitEthernet1/0/47,  Port ID (outgoing port): GigabitEthernet1/1/1
Holdtime : 156 sec

Version :
Cisco IOS Software [Cupertino], Catalyst L3 Switch Software (CAT9K_LITE_IOSXE), Version 17.9.4a

-------------------------
Device ID: AP-LOBBY-01
Entry address(es):
  IP address: 10.10.20.31
Platform: cisco C9120AXI-E,  Capabilities: Trans-Bridge Source-Route-Bridge IGMP
Interface: GigabitEthernet1/0/3,  Port ID (outgoing port): GigabitEthernet0
Holdtime : 122 sec

-------------------------
Device ID: SEP001122334455
Entry address(es):
  IP address: 10.10.30.5
Platform: Cisco IP Phone 8841,  Capabilities: Host Phone Two-port Mac Relay
Interface: GigabitEthernet1/0/5,  Port ID (outgoing port): Port 1
Holdtime : 130 sec
"""

LLDP_DETAIL = """------------------------------------------------
Chassis id: 0050.5686.aa01
Port id: eth0
Port Description: eth0
System Name: esx-host-07

System Description:
VMware ESXi 8.0.2 build-22380479

Time remaining: 96 seconds
System Capabilities: B
Enabled Capabilities: B
Management Addresses:
    IP: 10.10.40.17
Auto Negotiation - not supported
Local Intf: Te1/1/3

------------------------------------------------
Chassis id: 70:0f:6a:11:22:00
Port id: Gi1/0/47
Port Description: GigabitEthernet1/0/47
System Name: HQ-CORE-SW01

System Description:
Cisco IOS Software [Cupertino], Catalyst L3 Switch Software (CAT9K_IOSXE), Version 17.9.4a

System Capabilities: B,R
Enabled Capabilities: B,R
Management Addresses:
    IP: 10.10.0.1
Local Intf: Gi1/1/1
"""

MAC_TABLE_IOS = """          Mac Address Table
-------------------------------------------

Vlan    Mac Address       Type        Ports
----    -----------       --------    -----
 All    0100.0ccc.cccc    STATIC      CPU
  10    0050.5686.aa01    DYNAMIC     Te1/1/3
  10    a4bb.6d12.3456    DYNAMIC     Gi1/0/12
  20    001e.c9aa.bbcc    DYNAMIC     Gi1/0/47
  30    0011.2233.4455    DYNAMIC     Gi1/0/5
Total Mac Addresses for this criterion: 5
"""

MAC_TABLE_NXOS = """Legend:
        * - primary entry, G - Gateway MAC, (R) - Routed MAC, O - Overlay MAC
   VLAN     MAC Address      Type      age     Secure NTFY Ports
---------+-----------------+--------+---------+------+----+------------------
*   100     0050.5686.bb02   dynamic  0         F      F    Eth1/5
G     -     00de.fb11.2233   static   -         F      F    sup-eth1(R)
*   100     0050.5686.bb03   dynamic  0         F      F    Po10
"""

ARP_IOS = """Protocol  Address          Age (min)  Hardware Addr   Type   Interface
Internet  10.10.10.1              -   700f.6a11.2200  ARPA   Vlan10
Internet  10.10.10.50             3   a4bb.6d12.3456  ARPA   Vlan10
Internet  10.10.40.17            12   0050.5686.aa01  ARPA   Vlan40
"""

ARP_FORTIGATE = """Address           Age(min)   Hardware Addr      Interface
10.44.0.20        2          00:50:56:86:cc:01  internal
"""

FGT_STATUS = """Version: FortiGate-60F v7.0.12,build0523,230530 (GA.M)
Serial-Number: FGT60FTK20001234
Hostname: BR04-FGT-EDGE
Operation Mode: NAT
"""

APC_ABOUT = """E000: Success
Hardware Factory
---------------
Model Number: AP8953
Serial Number: 5A1234T00123
Hardware Revision: 05
Manufacture Date: 03/12/2021
MAC Address: 00 C0 B7 AA BB CC
Management Uptime: 12 Days 4 Hours 3 Minutes

Application Module
-----------------
Name: rpdu2g
Version: v6.9.6
"""

AP_CONFIG_GENERAL = """Cisco AP Name   : AP-LOBBY-01
MAC Address     : 4c77.6d00.0001
IP Address      : 10.10.20.31
Primary Controller IP Address : 10.10.5.10
AP Model        : C9120AXI-E
IOS Version     : 17.9.4.27
AP Serial Number: FGL2345A0AB
Cisco AP Name   : AP-DOCK-02
MAC Address     : 4c77.6d00.0002
IP Address      : 10.10.20.32
AP Model        : C9120AXI-E
IOS Version     : 17.9.4.27
AP Serial Number: FGL2345A0AC
"""

AP_SUMMARY = """Number of APs: 1

AP Name                            Slots    AP Model              Ethernet MAC    Location                          Country     IP Address                                 State
----------------------------------------------------------------------------------------------------------------------------------------------------------------------
AP-WAREHOUSE-01                    2        AIR-AP2802I-E-K9      00f2.8b00.0001  default location                  GB          10.60.20.5                                 Registered
"""


def _snapshot(*blocks: tuple[str, str]) -> str:
    return "\n".join(f"! ---- {command} ----\n{output}" for command, output in blocks)


# --- parsers ------------------------------------------------------------------------------


def test_command_matching_accepts_abbreviations_and_pipes():
    assert inv.command_matches("sh ver", "show version")
    assert inv.command_matches("show cdp neighbor detail", "show cdp neighbors detail")
    assert inv.command_matches("sh cdp nei det", "show cdp neighbors detail")
    assert inv.command_matches("show mac-address-table", "show mac address-table")
    assert inv.command_matches("show arp", "show ip arp")
    assert inv.command_matches("sh ap config general | i MAC Address|IP Address|AP Model", "show ap summary")
    assert not inv.command_matches("show cdp neighbor", "show cdp neighbors detail")
    assert not inv.command_matches("show run", "show version")


def test_show_version_ios_and_nxos():
    ios = inv.parse_show_version(IOS_SHOW_VERSION)
    assert ios["hostname"] == "HQ-CORE-SW01" and ios["model"] == "C9300-48P" and ios["serial"] == "FOC2345X0AB"
    assert ios["software"] == "17.9.4a" and ios["base_mac"] == "700f.6a11.2200" and ios["uptime"].startswith("1 year")
    nx = inv.parse_show_version(NXOS_SHOW_VERSION)
    assert nx["hostname"] == "DC1-LEAF-01" and nx["model"] == "Nexus9000 C93180YC-EX" and nx["serial"] == "FDO23456ABC"
    assert nx["software"] == "9.3(12)" and nx["uptime"].startswith("200 day")


def test_show_inventory_lists_every_component():
    items = inv.parse_show_inventory(SHOW_INVENTORY)
    assert [i["name"] for i in items] == ["c93xx Stack", "Switch 1", "Switch 1 - Power Supply A", "TenGigabitEthernet1/1/1"]
    assert items[2]["pid"] == "PWR-C1-715WAC-P" and items[2]["serial"] == "DTN2340V0AB" and items[3]["description"] == "SFP-10GBase-SR"


def test_cdp_and_lldp_detail_parsers():
    cdp = inv.parse_cdp_detail(CDP_DETAIL)
    assert len(cdp) == 3
    assert cdp[0] == {
        "name": "HQ-DIST-SW11.corp.local",
        "ip": "10.10.1.11",
        "platform": "cisco C9200L-48P-4X",
        "capabilities": "Switch IGMP",
        "local_port": "GigabitEthernet1/0/47",
        "remote_port": "GigabitEthernet1/1/1",
    }
    assert cdp[2]["platform"] == "Cisco IP Phone 8841" and cdp[2]["remote_port"] == "Port"  # 'Port 1' has a space; first token kept
    lldp = inv.parse_lldp_detail(LLDP_DETAIL)
    assert len(lldp) == 2
    assert lldp[0]["name"] == "esx-host-07" and lldp[0]["ip"] == "10.10.40.17" and lldp[0]["local_port"] == "Te1/1/3"
    assert lldp[0]["platform"].startswith("VMware ESXi") and lldp[0]["remote_port"] == "eth0"
    assert lldp[1]["name"] == "HQ-CORE-SW01" and lldp[1]["ip"] == "10.10.0.1"


def test_mac_and_arp_tables_across_platforms():
    ios = inv.parse_mac_table(MAC_TABLE_IOS)
    assert [(r["vlan"], r["mac"], r["port"]) for r in ios] == [
        ("10", "0050.5686.aa01", "Te1/1/3"),
        ("10", "a4bb.6d12.3456", "Gi1/0/12"),
        ("20", "001e.c9aa.bbcc", "Gi1/0/47"),
        ("30", "0011.2233.4455", "Gi1/0/5"),
    ]
    nx = inv.parse_mac_table(MAC_TABLE_NXOS)
    assert [(r["mac"], r["port"]) for r in nx] == [("0050.5686.bb02", "Eth1/5"), ("0050.5686.bb03", "Po10")]
    arp = inv.parse_arp(ARP_IOS)
    assert {r["ip"]: r["mac"] for r in arp} == {"10.10.10.1": "700f.6a11.2200", "10.10.10.50": "a4bb.6d12.3456", "10.10.40.17": "0050.5686.aa01"}
    fgt = inv.parse_arp(ARP_FORTIGATE)
    assert fgt == [{"ip": "10.44.0.20", "mac": "0050.5686.cc01", "interface": "internal", "age": "2"}]


def test_fortigate_apc_and_access_point_parsers():
    fgt = inv.parse_fortigate_status(FGT_STATUS)
    assert fgt["hostname"] == "BR04-FGT-EDGE" and fgt["model"] == "FortiGate-60F" and fgt["serial"] == "FGT60FTK20001234" and fgt["software"] == "v7.0.12,build0523,230530"
    apc = inv.parse_apc_about(APC_ABOUT)
    assert apc["model"] == "AP8953" and apc["serial"] == "5A1234T00123" and apc["base_mac"] == "00c0.b7aa.bbcc" and apc["software"] == "v6.9.6"
    aps = inv.parse_ap_config_general(AP_CONFIG_GENERAL)
    assert [(a["name"], a["model"], a["mac"], a["ip"], a["serial"]) for a in aps] == [
        ("AP-LOBBY-01", "C9120AXI-E", "4c77.6d00.0001", "10.10.20.31", "FGL2345A0AB"),
        ("AP-DOCK-02", "C9120AXI-E", "4c77.6d00.0002", "10.10.20.32", "FGL2345A0AC"),
    ]
    summary = inv.parse_ap_summary(AP_SUMMARY)
    assert summary == [{"name": "AP-WAREHOUSE-01", "model": "AIR-AP2802I-E-K9", "mac": "00f2.8b00.0001", "ip": "10.60.20.5", "serial": None, "software": None}]


# --- assembling ---------------------------------------------------------------------------


def _core_snapshot() -> str:
    return _snapshot(
        ("show ver", IOS_SHOW_VERSION),
        ("show inventory", SHOW_INVENTORY),
        ("show cdp neighbor detail", CDP_DETAIL),
        ("sh lldp nei detail", LLDP_DETAIL),
        ("show mac address-table", MAC_TABLE_IOS),
        ("show ip arp", ARP_IOS),
        ("show run", IOS_RUN_INTERFACES),
    )


def test_build_inventory_joins_everything_together():
    inputs = [
        inv.SnapshotInput("d1", "HQ-CORE-SW01", "10.10.0.1", "cisco_ios", "HQ", "core_switch", _core_snapshot()),
        inv.SnapshotInput("d2", "HQ-DIST-SW11", "10.10.1.11", "cisco_ios", "HQ", None, _snapshot(("show run", "hostname HQ-DIST-SW11"))),
        inv.SnapshotInput("d3", "DC1-LEAF-01", "10.20.0.11", "cisco_nxos", "DC1", None, _snapshot(("show version", NXOS_SHOW_VERSION), ("show mac address-table", MAC_TABLE_NXOS))),
        inv.SnapshotInput("d4", "NEVER-COLLECTED", "10.9.9.9", "cisco_ios", None, None, None),
    ]
    built = inv.build_inventory(inputs)

    core = built.devices[0]
    assert core.hostname == "HQ-CORE-SW01" and core.model == "C9300-48P" and core.serial == "FOC2345X0AB" and core.base_mac == "700f.6a11.2200"
    assert core.commands_missing == []
    dist = built.devices[1]
    assert dist.has_snapshot and dist.model is None and "show version" in dist.commands_missing and "show inventory" in dist.commands_missing
    assert built.devices[3].has_snapshot is False and built.devices[3].commands_missing == list(inv.INVENTORY_COMMANDS["cisco_ios"])

    assert len(built.hardware) == 4 and built.hardware[0].device_name == "HQ-CORE-SW01"

    # Neighbours: the dist switch is managed (matched by name, FQDN stripped); the AP, phone and ESX host are not.
    by_name = {n.name: n for n in built.neighbors}
    assert by_name["HQ-DIST-SW11.corp.local"].managed_device_name == "HQ-DIST-SW11"
    assert by_name["AP-LOBBY-01"].managed_device_id is None
    assert by_name["HQ-CORE-SW01"].managed_device_name == "HQ-CORE-SW01"  # LLDP sees the core itself via the dist uplink
    kinds = {u.name: u.kind for u in built.unmanaged}
    assert kinds == {"AP-LOBBY-01": "Access point", "SEP001122334455": "IP phone", "esx-host-07": "Host"}
    esx = next(u for u in built.unmanaged if u.name == "esx-host-07")
    assert esx.ip == "10.10.40.17" and esx.seen_from == ["HQ-CORE-SW01 Te1/1/3"] and esx.protocols == ["lldp"]

    # Endpoints: MAC table joined with ARP; ports with a neighbour are flagged as uplinks.
    endpoints = {(e.device_name, e.mac): e for e in built.endpoints}
    assert endpoints[("HQ-CORE-SW01", "a4bb.6d12.3456")].ip == "10.10.10.50" and endpoints[("HQ-CORE-SW01", "a4bb.6d12.3456")].on_uplink is False
    assert endpoints[("HQ-CORE-SW01", "001e.c9aa.bbcc")].on_uplink is True  # Gi1/0/47 leads to HQ-DIST-SW11
    assert endpoints[("HQ-CORE-SW01", "0050.5686.aa01")].on_uplink is False  # Te1/1/3 has the ESX host on LLDP: a host, not an uplink
    assert endpoints[("HQ-CORE-SW01", "0011.2233.4455")].on_uplink is False  # the phone's port likewise
    assert endpoints[("DC1-LEAF-01", "0050.5686.bb02")].ip is None
    assert built.models["C9300-48P"] == 1 and built.models["Nexus9000 C93180YC-EX"] == 1 and built.models["cisco C9120AXI-E"] == 1


def test_wlc_inventory_uses_sysinfo_or_version_and_lists_access_points():
    inputs = [
        inv.SnapshotInput("w1", "HQ-WLC-9800", "10.10.5.10", "cisco_wlc_9800", "HQ", None, _snapshot(("sh ap config general | i MAC Address|IP Address|AP Model", AP_CONFIG_GENERAL), ("show vers", IOS_SHOW_VERSION))),
        inv.SnapshotInput("w2", "PLANT-WLC", "10.60.5.10", "cisco_wlc", None, None, _snapshot(("show ap summary", AP_SUMMARY), ("show inventory", 'NAME: "Chassis", DESCR: "Cisco 3504 Wireless Controller"\nPID: AIR-CT3504-K9, VID: V01, SN: FCW2222A0AB\n'))),
    ]
    built = inv.build_inventory(inputs)
    assert [(a.controller_name, a.name) for a in built.access_points] == [("HQ-WLC-9800", "AP-LOBBY-01"), ("HQ-WLC-9800", "AP-DOCK-02"), ("PLANT-WLC", "AP-WAREHOUSE-01")]
    plant = built.devices[1]
    assert plant.serial == "FCW2222A0AB" and plant.model == "AIR-CT3504-K9"  # falls back to the chassis in show inventory
    assert built.models["C9120AXI-E"] == 2


# --- API + export ---------------------------------------------------------------------------


async def _register(client: AsyncClient, email: str) -> str:
    resp = await client.post("/api/auth/register", json={"org_name": "InventoryOrg", "email": email, "password": "password123"})
    assert resp.status_code == 201, resp.text
    return resp.json()["access_token"]


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


async def _collect_all(client: AsyncClient, token: str, monkeypatch, content_by_host: dict[str, str]) -> None:
    monkeypatch.setattr(tasks_module, "_attempt_collection", lambda device, *a, **k: content_by_host[device.host])
    job = await client.post("/api/jobs", headers=_auth(token), json={})
    assert job.status_code == 201, job.text
    detail = await client.get(f"/api/jobs/{job.json()['id']}", headers=_auth(token))
    assert detail.json()["status"] == "completed", detail.text


async def test_inventory_api_and_excel_export(client: AsyncClient, unique_email, monkeypatch):
    token = await _register(client, unique_email)
    await client.post("/api/credentials", headers=_auth(token), json={"name": "lab", "username": "admin", "password": "cisco123"})
    for name, host, kind in (("HQ-CORE-SW01", "10.10.0.1", "cisco_ios"), ("BR04-FGT-EDGE", "10.44.0.1", "fortinet")):
        assert (await client.post("/api/devices", headers=_auth(token), json={"name": name, "host": host, "device_type": kind, "site": "HQ"})).status_code == 201

    empty = await client.get("/api/inventory", headers=_auth(token))
    assert empty.status_code == 200 and empty.json()["summary"]["devices_with_config"] == 0
    assert empty.json()["devices"][0]["has_snapshot"] is False

    await _collect_all(
        client,
        token,
        monkeypatch,
        {"10.10.0.1": _core_snapshot(), "10.44.0.1": _snapshot(("get system status", FGT_STATUS), ("get system arp", ARP_FORTIGATE))},
    )
    resp = await client.get("/api/inventory", headers=_auth(token))
    assert resp.status_code == 200, resp.text
    body = resp.json()
    # HQ-DIST-SW11 is not a managed device in this org, so it counts as unmanaged here (4 with the AP, phone and ESX host);
    # its uplink port Gi1/0/47 is the only one excluded from the endpoint count (4 MAC rows -> 3 endpoints).
    subnet_count = body["summary"].pop("subnets")
    assert body["summary"] == {"devices": 2, "devices_with_config": 2, "devices_with_serial": 2, "hardware": 4, "access_points": 0, "neighbors": 5, "unmanaged": 4, "endpoints": 3}
    by_net = {s["network"]: s for s in body["subnets"]}
    assert subnet_count == len(by_net) >= 6
    assert by_net["10.10.10.0/24"]["vlan"] == "10" and by_net["10.10.10.0/24"]["addresses"][0]["device_name"] == "HQ-CORE-SW01"
    # The FortiGate collected no config here, so its LAN is only known from addresses in use: its own and one ARP entry.
    assert by_net["10.44.0.0/24"]["source"] == "seen" and by_net["10.44.0.0/24"]["hosts_seen"] == 2
    core = next(d for d in body["devices"] if d["name"] == "HQ-CORE-SW01")
    fgt = next(d for d in body["devices"] if d["name"] == "BR04-FGT-EDGE")
    assert core["serial"] == "FOC2345X0AB" and core["snapshot_id"] and core["collected_at"]
    assert fgt["model"] == "FortiGate-60F" and fgt["serial"] == "FGT60FTK20001234" and fgt["commands_missing"] == ["show running-config"]

    export = await client.get("/api/inventory/export.xlsx", headers=_auth(token))
    assert export.status_code == 200, export.text
    assert export.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="InventoryOrg-inventory-' in export.headers["content-disposition"]
    wb = load_workbook(BytesIO(export.content))
    assert wb.sheetnames == ["Summary", "Devices", "Hardware", "Neighbors", "Access points", "Unmanaged", "Subnets", "Endpoints"]
    subnets = wb["Subnets"]
    assert [c.value for c in subnets[1]][:3] == ["Network", "Mask", "VLAN"] and subnets.max_row == subnet_count + 1
    devices = wb["Devices"]
    header = [c.value for c in devices[1]]
    assert header[:8] == ["Name", "Hostname (reported)", "Management IP", "Type", "Site", "Role", "Model", "Serial"]
    rows = {row[0]: row for row in devices.iter_rows(min_row=2, values_only=True)}
    assert rows["HQ-CORE-SW01"][7] == "FOC2345X0AB" and rows["BR04-FGT-EDGE"][6] == "FortiGate-60F"
    assert wb["Hardware"].max_row == 5 and wb["Unmanaged"].max_row == 5  # header + 4 components / 4 unmanaged devices
    assert wb["Summary"]["A1"].value == "Organization" and wb["Summary"]["B1"].value == "InventoryOrg"


async def test_command_coverage_and_one_click_add(client: AsyncClient, unique_email):
    token = await _register(client, unique_email)
    await client.post("/api/devices", headers=_auth(token), json={"name": "sw1", "host": "192.0.2.1", "device_type": "cisco_ios"})
    coverage = (await client.get("/api/inventory/commands", headers=_auth(token))).json()
    ios = next(c for c in coverage if c["device_type"] == "cisco_ios")
    # The built-in IOS list already runs everything the inventory reads
    # (abbreviated: "show ver", "sh lldp nei detail", "show mac address"...).
    assert ios["device_count"] == 1 and ios["is_custom_profile"] is False and ios["missing"] == []
    nxos = next(c for c in coverage if c["device_type"] == "cisco_nxos")
    nxos_extra = [c for c in inv.INVENTORY_COMMANDS["cisco_nxos"] if c != "show running-config"]
    assert nxos["device_count"] == 0 and nxos["missing"] == nxos_extra  # its default is only show running-config, which counts

    added = await client.post("/api/inventory/commands/cisco_nxos", headers=_auth(token))
    assert added.status_code == 200, added.text
    commands = added.json()["commands"]
    assert commands == ["show running-config", *nxos_extra]  # existing kept first, nothing duplicated
    assert added.json()["is_custom"] is True
    after = (await client.get("/api/inventory/commands", headers=_auth(token))).json()
    assert next(c for c in after if c["device_type"] == "cisco_nxos")["missing"] == []
    assert next(c for c in after if c["device_type"] == "cisco_nxos")["is_custom_profile"] is True
    # Idempotent, and a partial custom list only gets what it lacks.
    again = await client.post("/api/inventory/commands/cisco_nxos", headers=_auth(token))
    assert again.json()["commands"] == commands
    await client.put("/api/command-profiles/cisco_ios", headers=_auth(token), json={"commands": ["show run", "sh ver"]})
    ios_added = await client.post("/api/inventory/commands/cisco_ios", headers=_auth(token))
    assert ios_added.json()["commands"] == ["show run", "sh ver", "show inventory", "show cdp neighbors detail", "show lldp neighbors detail", "show mac address-table", "show ip arp"]
    assert (await client.post("/api/inventory/commands/versa", headers=_auth(token))).status_code == 404

    member_email = f"member-{unique_email}"
    await client.post("/api/users", headers=_auth(token), json={"email": member_email, "password": "password123"})
    member = (await client.post("/api/auth/login", json={"email": member_email, "password": "password123"})).json()["access_token"]
    assert (await client.post("/api/inventory/commands/cisco_nxos", headers=_auth(member))).status_code == 403
    assert (await client.get("/api/inventory", headers=_auth(member))).status_code == 200


# --- subnets ------------------------------------------------------------------------------

IOS_RUN_INTERFACES = """hostname HQ-CORE-SW01
!
interface Loopback0
 ip address 10.255.0.1 255.255.255.255
!
interface Vlan10
 description Users - floor 1
 ip address 10.10.10.1 255.255.255.0
 ip address 10.10.11.1 255.255.255.0 secondary
!
interface Vlan40
 description Servers
 vrf forwarding PROD
 ip address 10.10.40.1 255.255.255.0
!
interface GigabitEthernet1/0/48
 description Uplink to FW
 no switchport
 ip address 192.168.100.2 255.255.255.252
!
interface GigabitEthernet1/0/47.200
 encapsulation dot1Q 200
 ip address 172.16.200.1 255.255.255.128
!
interface GigabitEthernet1/0/1
 switchport mode access
!
interface Vlan99
 no ip address
!
router ospf 1
 network 10.0.0.0 0.255.255.255 area 0
"""

NXOS_RUN_INTERFACES = """interface Vlan20
  description Storage
  vrf member STORAGE
  ip address 10.20.20.1/24
interface Ethernet1/49
  description to core
  ip address 10.20.255.1/31
interface mgmt0
  vrf member management
  ip address 10.20.0.11/24
"""

FORTIGATE_CONFIG = """config system global
    set hostname "BR04-FGT-EDGE"
end
config system interface
    edit "wan1"
        set vdom "root"
        set ip 203.0.113.10 255.255.255.248
        set alias "ISP"
    next
    edit "internal"
        set vdom "root"
        set ip 10.44.0.1 255.255.255.0
        set description "Branch LAN"
        config secondaryip
            edit 1
                set ip 10.44.1.1 255.255.255.0
            next
        end
    next
    edit "guest"
        set vdom "root"
        set ip 10.44.50.1 255.255.255.0
        set interface "internal"
        set vlanid 50
    next
end
config firewall policy
end
"""

AIREOS_RUN_CONFIG = """Interface Configuration

Interface Name................................... management
MAC Address...................................... 00:a2:ee:11:22:33
IP Address....................................... 10.60.5.10
IP Netmask....................................... 255.255.255.0
IP Gateway....................................... 10.60.5.1
VLAN............................................. 5
Active Physical Port............................. LAG (13)

Interface Name................................... staff-wireless
IP Address....................................... 10.60.120.2
IP Netmask....................................... 255.255.254.0
IP Gateway....................................... 10.60.120.1
VLAN............................................. 120

Interface Name................................... virtual
IP Address....................................... 192.0.2.1
"""


def test_cisco_interface_addresses_ios_and_nxos():
    rows = inv.parse_cisco_interface_addresses(IOS_RUN_INTERFACES)
    by_ip = {r["ip"]: r for r in rows}
    assert set(by_ip) == {"10.255.0.1", "10.10.10.1", "10.10.11.1", "10.10.40.1", "192.168.100.2", "172.16.200.1"}
    assert by_ip["10.10.10.1"] == {
        "interface": "Vlan10",
        "ip": "10.10.10.1",
        "prefix_len": 24,
        "network": "10.10.10.0/24",
        "description": "Users - floor 1",
        "vlan": "10",
        "vrf": None,
        "secondary": False,
    }
    assert by_ip["10.10.11.1"]["secondary"] is True and by_ip["10.10.11.1"]["vlan"] == "10"
    assert by_ip["10.10.40.1"]["vrf"] == "PROD"
    assert by_ip["192.168.100.2"]["network"] == "192.168.100.0/30" and by_ip["192.168.100.2"]["vlan"] is None
    assert by_ip["172.16.200.1"]["vlan"] == "200" and by_ip["172.16.200.1"]["network"] == "172.16.200.0/25"

    nx = {r["ip"]: r for r in inv.parse_cisco_interface_addresses(NXOS_RUN_INTERFACES)}
    assert nx["10.20.20.1"]["network"] == "10.20.20.0/24" and nx["10.20.20.1"]["vrf"] == "STORAGE" and nx["10.20.20.1"]["vlan"] == "20"
    assert nx["10.20.255.1"]["network"] == "10.20.255.0/31"
    assert nx["10.20.0.11"]["interface"] == "mgmt0" and nx["10.20.0.11"]["vrf"] == "management"


def test_fortigate_and_aireos_interface_addresses():
    fgt = {r["ip"]: r for r in inv.parse_fortigate_interface_addresses(FORTIGATE_CONFIG)}
    assert set(fgt) == {"203.0.113.10", "10.44.0.1", "10.44.1.1", "10.44.50.1"}
    assert fgt["203.0.113.10"]["description"] == "ISP" and fgt["203.0.113.10"]["network"] == "203.0.113.8/29" and fgt["203.0.113.10"]["vrf"] == "root"
    assert fgt["10.44.0.1"]["description"] == "Branch LAN" and fgt["10.44.0.1"]["secondary"] is False
    assert fgt["10.44.1.1"]["interface"] == "internal" and fgt["10.44.1.1"]["secondary"] is True
    assert fgt["10.44.50.1"]["vlan"] == "50"

    wlc = {r["ip"]: r for r in inv.parse_aireos_interface_addresses(AIREOS_RUN_CONFIG)}
    assert set(wlc) == {"10.60.5.10", "10.60.120.2"}  # the virtual interface has no netmask
    assert wlc["10.60.5.10"]["interface"] == "management" and wlc["10.60.5.10"]["vlan"] == "5" and wlc["10.60.5.10"]["network"] == "10.60.5.0/24"
    assert wlc["10.60.120.2"]["network"] == "10.60.120.0/23" and wlc["10.60.120.2"]["vlan"] == "120"

    # Dispatcher: FortiOS and AireOS are recognised from the text even for a custom type.
    assert inv.parse_interface_addresses(FORTIGATE_CONFIG, "custom_fw") == inv.parse_fortigate_interface_addresses(FORTIGATE_CONFIG)
    assert inv.parse_interface_addresses(AIREOS_RUN_CONFIG, "cisco_wlc") == inv.parse_aireos_interface_addresses(AIREOS_RUN_CONFIG)
    assert inv.parse_interface_addresses(IOS_RUN_INTERFACES, "cisco_ios")[1]["network"] == "10.10.10.0/24"


def test_build_inventory_lists_subnets_with_hosts_and_inferred_ranges():
    inputs = [
        inv.SnapshotInput("d1", "HQ-CORE-SW01", "10.10.0.1", "cisco_ios", "HQ", None, _snapshot(("show run", IOS_RUN_INTERFACES), ("show ip arp", ARP_IOS))),
        inv.SnapshotInput("d2", "HQ-DIST-SW11", "10.10.1.11", "cisco_ios", "HQ", None, _snapshot(("show running-config", "interface Vlan10\n ip address 10.10.10.2 255.255.255.0\n"))),
        inv.SnapshotInput("f1", "BR04-FGT-EDGE", "10.44.0.1", "fortinet", "BR04", None, _snapshot(("show full-configuration", FORTIGATE_CONFIG), ("get system arp", ARP_FORTIGATE))),
    ]
    built = inv.build_inventory(inputs)
    by_net = {s.network: s for s in built.subnets}

    users = by_net["10.10.10.0/24"]
    assert users.vlan == "10" and users.name == "Users - floor 1" and users.source == "config" and users.usable == 254
    assert [(a.device_name, a.ip) for a in users.addresses] == [("HQ-CORE-SW01", "10.10.10.1"), ("HQ-DIST-SW11", "10.10.10.2")]
    assert users.hosts_seen == 1  # 10.10.10.50 from ARP; the gateway's own ARP entry is an interface address, not a host
    assert by_net["10.10.40.0/24"].vrf == "PROD" and by_net["10.10.40.0/24"].hosts_seen == 1  # 10.10.40.17
    assert by_net["10.44.0.0/24"].hosts_seen == 1 and by_net["10.44.0.0/24"].name == "Branch LAN"  # FortiGate ARP
    assert by_net["192.168.100.0/30"].usable == 2 and by_net["10.255.0.1/32"].usable == 1

    # Management addresses of devices that sit in no configured subnet become inferred /24s.
    inferred = by_net["10.10.0.0/24"]
    assert inferred.source == "seen" and inferred.hosts_seen == 1 and inferred.addresses == [] and inferred.vlan is None
    assert by_net["10.10.1.0/24"].source == "seen"
    # Sorted by address.
    assert [s.network for s in built.subnets][:3] == ["10.10.0.0/24", "10.10.1.0/24", "10.10.10.0/24"]
    # The FortiGate's own management address is one of its interface addresses, so it is not counted as a host anywhere.
    assert all("10.44.0.1" not in [a.ip for a in s.addresses] or s.network == "10.44.0.0/24" for s in built.subnets)
    assert built.devices[0].commands_missing == [c for c in inv.INVENTORY_COMMANDS["cisco_ios"] if c not in ("show running-config", "show ip arp")]
