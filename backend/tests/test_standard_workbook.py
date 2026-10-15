"""Importing the configuration-standard workbook into template profiles,
and auditing controllers against those profiles."""

from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import Workbook

from app.api import wireless_audit as wa_api
from app.schemas.wireless_audit import TemplateProfile, WirelessStandard
from app.services import inventory as inv
from app.services import wireless_audit as wa
from app.services.config_template import check_template
from app.services.standard_workbook import import_workbook, parse_sheet

# A small workbook laid out like the real one: A commands, B variables, C
# notes, with headings, an example block, exec steps and regional choices.
XE_ROWS = [
    ("Note: Base configuration is for rebuilding the WLC if needed", None, None),
    ("Commands", "Variables", "Notes"),
    ("service password-encryption", None, None),
    ("hostname <<hostname>>", "<<hostname>>", None),
    ("enable secret <<enable-secret>>", "<<enable-secret>>", "https://vault.example/"),
    (None, None, None),
    ("parameter-map type webauth global", None, None),
    (" type webauth", None, None),
    (" virtual-ip ipv4 192.0.2.1", None, None),
    (None, None, None),
    ("Radius servers", None, None),
    ("radius server ISE1", None, None),
    (" address ipv4 10.1.1.1 auth-port 1812 acct-port 1813", None, None),
    (" key <<radius key>>", "<<radius key>>", None),
    (None, None, None),
    ("crypto key generate rsa general-keys modulus 2048", None, None),
    ("wireless profile policy PP_WLAN_1", "CorpWiFi profile", None),
    ("shutdown", None, None),
    (" vlan datagroup100", None, None),
    (" no shutdown", None, None),
    (None, None, None),
    ("In a mix Access points 2800/3800 and 9100AX 802.11ax should be disabled in all WLAN ID", None, None),
    (None, None, None),
    ("wlan <<profilename>>  <<wlanid>>  <<ssid>>", None, None),
    ("shutdown", None, None),
    ("no dot11ax", None, None),
    ("no shutdown", None, None),
    (None, None, None),
    ("###example###", None, None),
    ("wlan CorpWiFi 1 CorpWiFi", None, None),
    (None, None, None),
    ("wlan CorpWiFi 1 CorpWiFi", "CorpWiFi SSID", None),
    (" security ft", None, None),
    ("radio policy dot11 5ghz", None, None),
    ("wlan Scanners 2 Scanners", None, "only for sites where it is used"),
    (" security wpa psk set-key ascii 0 <<pre-shared key>>", None, None),
    (None, None, None),
    ("wireless tag policy APG_<<site-prefix>>-APGROUP01", None, None),
    (" wlan CorpWiFi policy PP_WLAN_1", None, None),
    ("wlan Guest policy PP_WLAN_16", None, None),
    (None, None, None),
    ("parameter-map type webauth global", None, None),
    ("webauth-http-enable", None, None),
    ("wireless mobility mac-address xxxx.xxxx.xxxx", None, None),
]

AIREOS_ROWS = [
    ("Commands", "Variables", "Notes"),
    ("config time ntp server 1 192.168.128.1", None, None),
    ("config sysname <<hostname>>", "<<hostname>>", None),
    ("config wlan create 1 CorpWiFi CorpWiFi", None, None),
    ("config wlan disable 1", None, None),
    ("config wlan create 2 LegacyNet2 LegacyNet2", "Please do not configure if CorpWiFi is present on site", None),
    ("y", None, None),
    ("config snmp community delete public", None, None),
    (None, None, None),
    ("NALA", None, None),
    ("config radius auth add 1 137.0.0.1 1812 ascii <<radius secret>>", None, None),
    (None, None, None),
    ("Europe", None, None),
    ("config radius auth add 1 143.0.0.1 1812 ascii <<radius secret>>", None, None),
    (None, None, None),
    ("config network telnet disable", None, None),
    ("transfer download start", None, None),
    ("save config", None, None),
]


def workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "WLC 9800 IOS-XE"
    for row in XE_ROWS:
        ws.append(list(row))
    flex = wb.create_sheet("Flex vWLC 9800 IOS-XE")
    for row in XE_ROWS[:6]:
        flex.append(list(row))
    flex.append(["wireless profile flex <<site-prefix>>_FLEXPROFILE01", None, None])
    flex.append([" native-vlan-id 101", None, None])
    ae = wb.create_sheet("WLC AireOS")
    for row in AIREOS_ROWS:
        ae.append(list(row))
    wb.create_sheet("Empty")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sheets() -> dict:
    return {s.name: s for s in import_workbook(workbook_bytes(), "standard.xlsx").sheets}


def test_sheets_platform_and_kind():
    by_name = sheets()
    assert set(by_name) == {"WLC 9800 IOS-XE", "Flex vWLC 9800 IOS-XE", "WLC AireOS"}
    assert (by_name["WLC 9800 IOS-XE"].platform, by_name["WLC 9800 IOS-XE"].kind) == ("9800", "local")
    assert (by_name["Flex vWLC 9800 IOS-XE"].platform, by_name["Flex vWLC 9800 IOS-XE"].kind) == ("9800", "flex")
    assert by_name["WLC AireOS"].platform == "aireos"


def test_xe_template_shape():
    sheet = sheets()["WLC 9800 IOS-XE"]
    lines = [line for line in sheet.template.splitlines() if not line.startswith("#")]
    text = "\n".join(lines)
    # Placeholders, secrets, fillers.
    assert "hostname {{hostname}}" in lines
    assert "enable secret {{enable_secret}}" in lines
    assert " key {{radius_key}}" in lines
    assert "wireless mobility mac-address {{value}}" in lines
    # "ascii 0 <psk>" is stored encrypted ("ascii 8 ..."): the type digit is a wildcard.
    assert " security wpa psk set-key ascii {{enc}} {{pre_shared_key}}" in lines
    # Headings, exec, examples and procedural shutdowns dropped.
    assert "Radius servers" not in text and "crypto key" not in text
    assert "shutdown\n" not in text.replace("no shutdown", "")
    # Flush-left children nested under their block.
    assert "wireless profile policy PP_WLAN_1\n vlan datagroup100\n no shutdown" in text
    assert "wlan CorpWiFi 1 CorpWiFi\n security ft\n radio policy dot11 5ghz" in text
    assert "wireless tag policy APG_{{site_prefix}}-APGROUP01\n wlan CorpWiFi policy PP_WLAN_1\n wlan Guest policy PP_WLAN_16" in text
    # The conditional heading makes the wildcard WLAN block optional.
    assert "optional: wlan {{profilename}} {{wlanid}} {{ssid}}\n optional: no dot11ax" in text
    # Notes saying "only for sites where it is used" make a block optional.
    assert "optional: wlan Scanners 2 Scanners" in text
    # A block configured in two places is merged into one.
    assert text.count("parameter-map type webauth global") == 1
    assert "parameter-map type webauth global\n type webauth\n virtual-ip ipv4 192.0.2.1\n webauth-http-enable" in text
    # Only one CorpWiFi WLAN block (the example's copy is ignored).
    assert text.count("wlan CorpWiFi 1 CorpWiFi") == 1
    reasons = {x.reason for x in sheet.skipped}
    assert {"heading or note", "example block", "exec command, not part of the running configuration"} <= reasons
    assert any(r.startswith("procedural step") for r in reasons)


def test_aireos_template_shape():
    sheet = sheets()["WLC AireOS"]
    lines = [line for line in sheet.template.splitlines() if not line.startswith("#")]
    assert "config sysname {{hostname}}" in lines
    assert "config wlan disable 1" not in lines
    assert "optional: config wlan create 2 LegacyNet2 LegacyNet2" in lines
    assert "absent: config snmp community create public {{*}}" in lines
    # Regional alternatives are optional; a blank row ends the region.
    assert "optional: config radius auth add 1 137.0.0.1 1812 ascii {{radius_secret}}" in lines
    assert "optional: config radius auth add 1 143.0.0.1 1812 ascii {{radius_secret}}" in lines
    assert "config network telnet disable" in lines
    assert not any(line.startswith(("transfer", "save", "y")) for line in lines)


def test_imported_template_matches_a_real_running_config():
    template = sheets()["WLC 9800 IOS-XE"].template
    config = """service password-encryption
hostname HQ-WLC01
enable secret 9 $9$abcdef
parameter-map type webauth global
 type webauth
 virtual-ip ipv4 192.0.2.1
 webauth-http-enable
radius server ISE1
 address ipv4 10.1.1.1 auth-port 1812 acct-port 1813
 key 7 08224F4B0A
wireless profile policy PP_WLAN_1
 vlan datagroup100
 no shutdown
wlan CorpWiFi 1 CorpWiFi
 no dot11ax
 security ft
 radio policy dot11 5ghz
 no shutdown
wireless tag policy APG_HQ01-APGROUP01
 wlan CorpWiFi policy PP_WLAN_1
 wlan Guest policy PP_WLAN_16
wireless mobility mac-address 0050.56bf.9071
"""
    result = check_template(template, config)
    failures = [i for i in result.issues if i.severity == "fail"]
    assert failures == [], failures
    # Optional Scanners WLAN is missing: a warning only.
    assert any("Scanners" in i.rule for i in result.issues)


def test_parse_sheet_ignores_stray_columns_and_blank_sheets():
    sheet = parse_sheet("Custom", [(1, "logging host 10.0.0.1", ""), (2, "Some heading", "")])
    assert sheet.platform == "9800"
    assert "logging host 10.0.0.1" in sheet.template
    assert "Some heading" not in sheet.template


# --- profile selection in the audit ------------------------------------------------------


def controller(model: str | None, config: str, name="WLC", device_type="cisco_wlc_9800") -> wa.ControllerFacts:
    f = wa.ControllerFacts(input=inv.SnapshotInput("id-" + name, name, "10.0.0.1", device_type, None, None, "x"), platform=wa.PLATFORMS[device_type])
    f.model = model
    f.config = config
    from app.services import wlc_parsers as wp

    f.tags = wp.parse_tag_config(config)
    return f


PROFILES = [
    TemplateProfile(name="WLC 9800 IOS-XE", platform="9800", kind="local", template="hostname {{x}}\n"),
    TemplateProfile(name="Flex vWLC 9800 IOS-XE", platform="9800", kind="flex", template="hostname {{x}}\n"),
    TemplateProfile(name="WLC AireOS", platform="aireos", kind="local", template="config sysname {{x}}\n"),
]


def test_select_profile():
    s = WirelessStandard(template_profiles=PROFILES)
    assert wa.select_profile(controller("C9800-40-K9", "hostname A"), s).name == "WLC 9800 IOS-XE"
    assert wa.select_profile(controller("C9800-CL-K9", "hostname A"), s).name == "Flex vWLC 9800 IOS-XE"
    remote = "wireless tag site ST-1\n flex-profile F1\n no local-site\n"
    assert wa.select_profile(controller("C9800-40-K9", remote), s).name == "Flex vWLC 9800 IOS-XE"
    assert wa.select_profile(controller(None, "", device_type="cisco_wlc"), s).name == "WLC AireOS"
    s.profile_assignments = {"WLC": "WLC 9800 IOS-XE"}
    assert wa.select_profile(controller("C9800-CL-K9", "hostname A"), s).name == "WLC 9800 IOS-XE"
    s.profile_assignments = {"WLC": "none"}
    assert wa.select_profile(controller("C9800-CL-K9", "hostname A"), s) is None


# --- API -----------------------------------------------------------------------------------


async def _register(client: AsyncClient, email: str) -> dict:
    resp = await client.post("/api/auth/register", json={"email": email, "password": "supersecret1", "org_name": "Standards Org"})
    assert resp.status_code in (200, 201), resp.text
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


@pytest.mark.asyncio
async def test_import_api_replaces_profiles_and_keeps_settings(client: AsyncClient, unique_email, monkeypatch):
    headers = await _register(client, unique_email)
    resp = await client.put("/api/wireless-audit/standard", json=WirelessStandard(ntp_servers=["10.0.0.123"]).model_dump(), headers=headers)
    assert resp.status_code == 200

    files = {"file": ("standard-v1.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    resp = await client.post("/api/wireless-audit/standard/import", files=files, headers=headers)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert {p["name"] for p in body["profiles"]} == {"WLC 9800 IOS-XE", "Flex vWLC 9800 IOS-XE", "WLC AireOS"}
    assert body["replaced"] == []
    std = body["standard"]["standard"]
    assert std["ntp_servers"] == ["10.0.0.123"]
    assert len(std["import_history"]) == 1 and std["import_history"][0]["filename"] == "standard-v1.xlsx"

    files = {"file": ("standard-v2.xlsx", workbook_bytes(), "application/octet-stream")}
    resp = await client.post("/api/wireless-audit/standard/import", files=files, headers=headers)
    assert resp.status_code == 200
    assert set(resp.json()["replaced"]) == {"WLC 9800 IOS-XE", "Flex vWLC 9800 IOS-XE", "WLC AireOS"}
    std = resp.json()["standard"]["standard"]
    assert len(std["template_profiles"]) == 3
    assert [h["filename"] for h in std["import_history"]] == ["standard-v2.xlsx", "standard-v1.xlsx"]

    resp = await client.post("/api/wireless-audit/standard/import", files={"file": ("bad.xlsx", b"not a workbook", "application/octet-stream")}, headers=headers)
    assert resp.status_code == 400

    # The audit then checks against the imported profile.
    async def fake_inputs(db, org_id):
        return [inv.SnapshotInput("x", "HQ-WLC01", "10.0.0.5", "cisco_wlc_9800", None, None, "! ---- show running-config ----\nhostname HQ-WLC01\n", None, "s")]

    monkeypatch.setattr(wa_api, "_inputs", fake_inputs)
    resp = await client.get("/api/wireless-audit?dns=false", headers=headers)
    assert resp.status_code == 200, resp.text
    c = resp.json()["controllers"][0]
    assert c["template_profile"] == "WLC 9800 IOS-XE"
    assert any(ch["check_id"] == "cfg.template" and ch["status"] == "fail" for ch in c["checks"])
