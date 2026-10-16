"""Wireless Audit: parsers for the extra controller outputs, the golden
template matcher, every check against sample 9800 and AireOS snapshots,
and the API (standard, audit, export, command coverage)."""

from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.api import wireless_audit as wa_api
from app.schemas.wireless_audit import WirelessStandard
from app.services import config_template as ct
from app.services import inventory as inv
from app.services import wireless_audit as wa
from app.services import wlc_parsers as wp

NOW = datetime(2026, 10, 15, 12, 0, tzinfo=timezone.utc)


def snapshot(sections: dict[str, str]) -> str:
    return "\n".join(f"! ---- {cmd} ----\n{out}" for cmd, out in sections.items())


# --- sample 9800 output ----------------------------------------------------------------------

XE_VERSION = """Cisco IOS XE Software, Version 17.09.05
Cisco IOS Software [Cupertino], C9800 Software (C9800_IOSXE-K9), Version 17.9.5, RELEASE SOFTWARE (fc2)
HQ-WLC01 uptime is 20 weeks, 1 day, 3 hours, 10 minutes
cisco C9800-40-K9 (KATAR) processor (revision KATAR) with 4167632K/6147K bytes of memory.
Processor board ID TTM2345ABCD
"""

XE_AP_GENERAL = """Cisco AP Name   : HQ-B01-F01-AP001
MAC Address                                     : a4b2.3c4d.5e60
IP Address                                      : 10.20.0.11
AP Model                                        : C9130AXI-B
IOS Version                                     : 17.9.5.47
AP Serial Number                                : FJC1111AAAA
Primary Cisco Controller Name                   : HQ-WLC01
Cisco AP Name   : ap-lobby
MAC Address                                     : a4b2.3c4d.5e70
IP Address                                      : 10.99.0.12
AP Model                                        : AIR-AP1832I-B-K9
IOS Version                                     : 17.6.1.13
AP Serial Number                                : FJC2222BBBB
Primary Cisco Controller Name                   : Not Configured
Cisco AP Name   : a4b2.3c4d.5e80
MAC Address                                     : a4b2.3c4d.5e80
IP Address                                      : 10.20.0.11
AP Model                                        : C9130AXI-B
IOS Version                                     : 17.9.5.47
AP Serial Number                                : FJC3333CCCC
Primary Cisco Controller Name                   : HQ-WLC01
"""

XE_AP_SUMMARY = """Number of APs: 3

AP Name                          Slots    AP Model             Ethernet MAC    Radio MAC       Location                          Country     IP Address                                 State
----------------------------------------------------------------------------------------------------------------------------------------------------------------------------------------------
HQ-B01-F01-AP001                 2        C9130AXI-B           a4b2.3c4d.5e60  a4b2.3c4d.0060  default location                  US          10.20.0.11                                 Registered
ap-lobby                         2        AIR-AP1832I-B-K9     a4b2.3c4d.5e70  a4b2.3c4d.0070  Lobby                             CA          10.99.0.12                                 Registered
a4b2.3c4d.5e80                   2        C9130AXI-B           a4b2.3c4d.5e80  a4b2.3c4d.0080  default location                  US          10.20.0.11                                 Downloading
"""

XE_TAGS = """Number of APs: 3

AP Name                           AP Mac           Site Tag Name                     Policy Tag Name                   RF Tag Name                       Misconfigured    Tag Source
----------------------------------------------------------------------------------------------------------------------------------------------------------------------------------------------
HQ-B01-F01-AP001                  a4b2.3c4d.5e60   st-hq                             pt-corp                           rf-hq                             No               Static
ap-lobby                          a4b2.3c4d.5e70   default-site-tag                  default-policy-tag                default-rf-tag                    No               Default
a4b2.3c4d.5e80                    a4b2.3c4d.5e80   st-hq                             pt-corp                           rf-hq                             Yes              Static
"""

XE_ETHERNET = """Ethernet Stats

AP Name                 Ethernet  Status     Speed               Rx Packets  Tx Packets  Discarded Packets
----------------------------------------------------------------------------------------------------------
HQ-B01-F01-AP001        Gi0       Up         2500 Mbps Full Duplex   12345      23456       0
ap-lobby                Gi0       Up         100 Mbps Half Duplex    1234       2345        12
"""

XE_5GHZ = """Number of 802.11a radios: 3

AP Name                           Mac Address     Slot    Admin State    Oper State    Width  Txpwr           Channel                                Mode
---------------------------------------------------------------------------------------------------------------------------------------------------------------
HQ-B01-F01-AP001                  a4b2.3c4d.0060  1       Enabled        Up            20     *2/8 (17 dBm)   (36)*                                  Local
ap-lobby                          a4b2.3c4d.0070  1       Enabled        Down          20     *1/8 (20 dBm)   (149)*                                 Local
a4b2.3c4d.5e80                    a4b2.3c4d.0080  1       Disabled       Down          20     *1/8 (20 dBm)   (40)*                                  Local
"""

XE_24GHZ = """AP Name                           Mac Address     Slot    Admin State    Oper State    Width  Txpwr           Channel                                Mode
---------------------------------------------------------------------------------------------------------------------------------------------------------------
HQ-B01-F01-AP001                  a4b2.3c4d.0060  0       Enabled        Up            20     *3/8 (11 dBm)   (6)*                                   Local
ap-lobby                          a4b2.3c4d.0070  0       Enabled        Up            20     *3/8 (11 dBm)   (3)*                                   Local
"""

XE_WLANS = """Number of WLANs: 3

ID   Profile Name                     SSID                             Status Security
------------------------------------------------------------------------------------------------------------------------
1    corp                             CORP                             UP     [WPA2][802.1x][AES]
2    guest                            GUEST                            UP     [open],MAC Filtering
3    legacy                           LEGACY                           DOWN   [WPA][PSK][TKIP]
"""

XE_CLIENTS = """Number of Clients: 3

MAC Address    AP Name                                        Type ID   State             Protocol Method     Role
-------------------------------------------------------------------------------------------------------------------------
1111.2222.3333 HQ-B01-F01-AP001                               WLAN 1    Run               11ax(5)  Dot1x      Local
4444.5555.6666 HQ-B01-F01-AP001                               WLAN 1    IP Learn          11n(2.4) Dot1x      Local
7777.8888.9999 ap-lobby                                       WLAN 2    Run               11g      None       Local
"""

XE_TRACKING = """vlan  Network Layer Address  Link Layer Address  Interface  mode  Time left
20    10.30.0.50             1111.2222.3333      Ca1        DHCP  1000 s
20    169.254.10.10          7777.8888.9999      Ca1        ARP   500 s
"""


def xe_client_detail(rssi: int, snr: int, state: str = "Run") -> str:
    return f"""Client MAC Address : 1111.2222.3333
Client Username : alice
AP Name: HQ-B01-F01-AP001
Client State : Associated
Policy Manager State: {state}
Authentication Key Management : FT-802.1x
Radio Signal Strength Indicator : {rssi} dBm
Signal to Noise Ratio : {snr} dB
"""


XE_RUN = """hostname HQ-WLC01
enable secret 9 $9$abc
service password-encryption
aaa new-model
ip http server
ip http secure-server
ip name-server 10.0.0.53
ntp server 10.0.0.123
logging host 10.0.0.514
snmp-server community public RO
wlan corp 1 CORP
 security dot1x authentication-list ISE
 no shutdown
wlan guest 2 GUEST
 no security wpa
 no shutdown
wlan legacy 3 LEGACY
 security wpa wpa1
 security wpa wpa2 ciphers tkip
line vty 0 4
 transport input ssh
line vty 5 15
 transport input telnet ssh
"""

XE_REDUNDANCY = """Redundant System Information :
       Available system uptime = 20 weeks, 1 day
Switchovers system experienced = 0
              Standby failures = 0
        Last switchover reason = none

                 Hardware Mode = Duplex
    Configured Redundancy Mode = sso
     Operating Redundancy Mode = sso
              Maintenance Mode = Disabled
                Communications = Up

Current Processor Information :
               Active Location = slot 1
        Current Software state = ACTIVE
Peer Processor Information :
              Standby Location = slot 2
        Current Software state = STANDBY HOT
"""

XE_NTP = """  address         ref clock       st   when   poll reach  delay  offset   disp
*~10.0.0.123      .GPS.            1     30     64   377  1.000   0.100  0.200
 * sys.peer, # selected, + candidate, - outlyer, x falseticker, ~ configured
"""

XE_LOG = """*Oct 15 10:00:00.123: %CAPWAPAC_SMGR-4-AP_DISJOIN: Chassis 1 R0/0: wncd: AP ap-lobby disjoined, reason: heartbeat timeout
*Oct 15 10:05:00.123: %CAPWAPAC_SMGR-4-AP_DISJOIN: Chassis 1 R0/0: wncd: AP ap-lobby disjoined, reason: heartbeat timeout
*Oct 15 10:06:00.123: %SYS-2-MALLOCFAIL: Memory allocation of 1024 bytes failed
*Oct 15 10:07:00.123: %SYS-5-CONFIG_I: Configured from console by admin
*Oct 15 10:08:00.123: %DOT1X-3-SERVER_DEAD: RADIUS server 10.0.0.10 is dead
"""


XE_TAG_CONFIG = """wireless tag site st-hq
 ap-profile APG_HQ
 description "HQ"
wireless tag site HQ01-APGROUP01-Flex
 flex-profile HQ01_FLEXPROFILE01
 no local-site
wireless tag policy pt-corp
 wlan CORP policy PP_WLAN_1
wireless tag rf rf-hq
 24ghz-rf-policy RF24_HQ
 5ghz-rf-policy RF5_HQ
"""

XE_AFFINITY = """Tag                              Type                No of AP's Joined   Load Factor   Wncd Instance
---------------------------------------------------------------------------------------------------------
st-hq                            SITE TAG            90                  0.18          0
HQ01-APGROUP01-Flex              SITE TAG            2                   0.01          1
default-site-tag                 SITE TAG            1                   0.00          2
"""

XE_WNCD_CPU = """  19828    19815   10%    9%    9%   S         452484 wncd_0
  19900    19815    1%    1%   85%   S         452484 wncd_1
  19950    19815    0%    0%    0%   S         452484 wncd_2
"""


def xe_snapshot() -> str:
    return snapshot(
        {
            "term len 0": "",
            "sh ap config general | i MAC Address|IP Address|AP Model|IOS Version|AP Serial Number|Cisco AP Name|Primary Cisco Controller Name": XE_AP_GENERAL,
            "show vers": XE_VERSION,
            "show ap summary": XE_AP_SUMMARY,
            "show ap ethernet statistics": XE_ETHERNET,
            "show wlan summary": XE_WLANS,
            "show running-config": XE_RUN + XE_TAG_CONFIG,
            "show wireless client summary": XE_CLIENTS,
            "show wireless device-tracking database ip": XE_TRACKING,
            "show wireless client mac-address 1111.2222.3333 detail": xe_client_detail(-80, 12),
            "show ap tag summary": XE_TAGS,
            "show ap dot11 24ghz summary": XE_24GHZ,
            "show ap dot11 5ghz summary": XE_5GHZ,
            "show redundancy": XE_REDUNDANCY,
            "show ntp associations": XE_NTP,
            "show processes cpu sorted | i CPU utilization": "CPU utilization for five seconds: 95%/1%; one minute: 91%; five minutes: 88%",
            "show logging": XE_LOG,
            "show wireless loadbalance tag affinity": XE_AFFINITY,
            "show processes cpu platform sorted | i wncd": XE_WNCD_CPU,
        }
    )


# --- sample AireOS output --------------------------------------------------------------------

AIREOS_SYSINFO = """Manufacturer's Name.............................. Cisco Systems Inc.
Product Name..................................... Cisco Controller
Product Version.................................. 8.10.185.0
System Name...................................... BR-WLC01
"""

AIREOS_AP_SUMMARY = """Number of APs.................................... 2

AP Name             Slots  AP Model              Ethernet MAC       Location          Country     IP Address       Clients   DSE Location
------------------  -----  --------------------  -----------------  ----------------  ----------  ---------------  --------  --------------
BR1-B01-F02-AP010    2     AIR-AP2802I-B-K9      00:11:22:33:44:10  2nd floor         US          10.40.0.10        60       [0 ,0 ,0 ]
BR1-B01-F02-AP011    2     AIR-AP2802I-B-K9      00:11:22:33:44:11  2nd floor         US          10.40.0.11        1        [0 ,0 ,0 ]
"""

AIREOS_RADIOS_A = """AP Name             MAC Address        Slot  Admin State  Oper State  Channel  TxPower
------------------  -----------------  ----  -----------  ----------  -------  -------
BR1-B01-F02-AP010   00:11:22:33:44:a0  1     ENABLED      UP          (36)*    *1/7 (17 dBm)
BR1-B01-F02-AP011   00:11:22:33:44:a1  1     ENABLED      UP          (44)*    *2/7 (14 dBm)
"""

AIREOS_RADIOS_B = """AP Name             MAC Address        Slot  Admin State  Oper State  Channel  TxPower
------------------  -----------------  ----  -----------  ----------  -------  -------
BR1-B01-F02-AP010   00:11:22:33:44:b0  0     ENABLED      UP          11*      *3/8 (11 dBm)
BR1-B01-F02-AP011   00:11:22:33:44:b1  0     DISABLED     DOWN        1*       *3/8 (11 dBm)
"""

AIREOS_WLANS = """Number of WLANs.................................. 2

WLAN ID  WLAN Profile Name / SSID                       Status    Interface Name        PMIPv6 Mobility
-------  -------------------------------------          --------  --------------------  ---------------
1        corp / CORP                                    Enabled   management            none
2        guest / GUEST                                  Enabled   guest                 none
"""

AIREOS_CLIENTS = """Number of Clients................................ 2

MAC Address       AP Name                        Slot Status        WLAN  Auth Protocol         Port Wired Tunnel  Role
----------------- ------------------------------ ---- ------------- ----- ---- ---------------- ---- ----- ------- ----------------
a4:bb:6d:12:34:56 BR1-B01-F02-AP010              1    Associated    1     Yes  802.11ac(5 GHz)  1    No    No      Local
a4:bb:6d:12:34:57 BR1-B01-F02-AP010              1    Associated    1     Yes  802.11ac(5 GHz)  1    No    No      Local
"""

AIREOS_CLIENT_DETAIL = """Client MAC Address............................... a4:bb:6d:12:34:56
Client Username ................................. bob
AP Name.......................................... BR1-B01-F02-AP010
Client State..................................... Associated
Policy Manager State............................. RUN
      Radio Signal Strength Indicator............ -55 dBm
      Signal to Noise Ratio...................... 40 dB
"""

AIREOS_CLIENT_DETAIL_2 = """Client MAC Address............................... a4:bb:6d:12:34:57
Policy Manager State............................. 8021X_REQD
      Radio Signal Strength Indicator............ -70 dBm
      Signal to Noise Ratio...................... 22 dB
"""

AIREOS_RUN_COMMANDS = """config time ntp server 1 10.0.0.123
config logging syslog host 10.0.0.99
config wlan security wpa disable 2
config wlan security web-auth enable 2
config network telnet enable
config 802.11a disable network
"""

AIREOS_NETWORK = """RF-Network Name............................. corp
Web Mode.................................... Disable
Secure Web Mode............................. Enable
Telnet...................................... Enable
"""

AIREOS_TIME = """Time............................................. Thu Oct 15 11:00:00 2026
NTP Servers
    NTP Polling Interval.........................     3600
     Index     NTP Key Index      NTP Server     Status    NTP Msg Auth Status
    -------  ---------------------------------------------------------------
       1            0             10.0.0.123     Not Synched  AUTH DISABLED
"""

AIREOS_REDUNDANCY = """ Redundancy Mode = SSO DISABLED
     Local State = ACTIVE
      Peer State = N/A
"""

AIREOS_TRAPLOG = """Number of Traps Since Last Reset ............ 3
Number of Traps Since Log Last Displayed .... 3

Log System Time              Trap
--- ------------------------ -------------------------------------------------
  0 Thu Oct 15 10:00:00 2026 AP Disassociated. Base Radio MAC:00:11:22:33:44:a0 ApName - BR1-B01-F02-AP010
  1 Thu Oct 15 09:00:00 2026 AP Disassociated. Base Radio MAC:00:11:22:33:44:a1 ApName - BR1-B01-F02-AP011
  2 Thu Oct 15 08:00:00 2026 Client Association: Client MAC:a4:bb:6d:12:34:56 Base Radio MAC:00:11:22:33:44:a0
"""


AIREOS_APGROUPS = """Total Number of AP Groups........................ 2

Site Name........................................ BR1-APGROUP01
Site Description................................. Branch office

RF Profile
----------
2.4 GHz band..................................... RF24_BR
5 GHz band....................................... <none>

WLAN ID          Interface          Network Admission Control          Radio Policy
-------          -----------        --------------------------         ------------
 1               management           Disabled                          All
 2               guest                Disabled                          All

AP Name             Slots  AP Model              Ethernet MAC       Location          Port  Country  Priority
------------------  -----  --------------------  -----------------  ----------------  ----  -------  --------
BR1-B01-F02-AP010    2     AIR-AP2802I-B-K9      00:11:22:33:44:10  2nd floor         1     US       1
BR1-B01-F02-AP012    2     C9120AXI-B            00:11:22:33:44:12  2nd floor         1     US       1

Site Name........................................ default-group
Site Description................................. <none>

AP Name             Slots  AP Model              Ethernet MAC       Location          Port  Country  Priority
------------------  -----  --------------------  -----------------  ----------------  ----  -------  --------
BR1-B01-F02-AP011    2     AIR-AP2802I-B-K9      00:11:22:33:44:11  2nd floor         1     US       1
"""


def aireos_snapshot() -> str:
    return snapshot(
        {
            "show ap summary": AIREOS_AP_SUMMARY,
            "show advanced 802.11a summary": AIREOS_RADIOS_A,
            "show advanced 802.11b summary": AIREOS_RADIOS_B,
            "show wlan apgroups": AIREOS_APGROUPS,
            "show sysinfo": AIREOS_SYSINFO,
            "show redundancy summary": AIREOS_REDUNDANCY,
            "show run-config": "System Inventory\n...",
            "show wlan summary": AIREOS_WLANS,
            "show client summary": AIREOS_CLIENTS,
            "show client detail a4:bb:6d:12:34:56": AIREOS_CLIENT_DETAIL,
            "show client detail a4:bb:6d:12:34:57": AIREOS_CLIENT_DETAIL_2,
            "show run-config commands": AIREOS_RUN_COMMANDS,
            "show network summary": AIREOS_NETWORK,
            "show time": AIREOS_TIME,
            "show cpu": "Current CPU(s) load: 5%\nIndividual CPU load: 5%/0%",
            "show msglog": "*apfReceiveTask: Oct 15 10:00:00.123: %APF-3-AUTHENTICATION_TRAP: Client failed to authenticate",
            "show traplog": AIREOS_TRAPLOG,
        }
    )


def inputs(xe_age_hours: float = 1) -> list[inv.SnapshotInput]:
    return [
        inv.SnapshotInput("wlc-xe", "HQ-WLC01", "10.0.0.5", "cisco_wlc_9800", "HQ", "wlc", xe_snapshot(), NOW - timedelta(hours=xe_age_hours), "snap-xe"),
        inv.SnapshotInput("wlc-ae", "BR-WLC01", "br-wlc01.corp.example", "cisco_wlc", "Branch", "wlc", aireos_snapshot(), NOW - timedelta(hours=1), "snap-ae"),
        inv.SnapshotInput("wlc-new", "NEW-WLC", "10.0.0.6", "cisco_wlc_9800", None, None, None),
    ]


STANDARD = WirelessStandard(
    ap_name_pattern=r"^[A-Z0-9]{2,3}-B\d{2}-F\d{2}-AP\d{3}$",
    ap_name_case="upper",
    dns_domain="corp.example",
    ap_subnets=["10.20.0.0/22", "10.40.0.0/22"],
    approved_ap_models=["C9130", "AIR-AP2802"],
    ap_group_name_pattern=r"^[A-Z0-9]{2,8}-APGROUP\d{2}(-(Production|Warehouse|Flex))?(-Outside)?$",
    country_codes=["US"],
    approved_9800_versions=["17.9.6", "17.12"],
    approved_aireos_versions=["8.10.185"],
    ntp_servers=["10.0.0.123"],
    syslog_servers=["10.0.0.514", "10.0.0.99"],
    dns_servers=["10.0.0.53"],
    require_ha=True,
    required_ssids=["CORP", "VOICE"],
    open_ssids_allowed=[],
    template_9800="hostname {{device_name}}\nline vty {{*}}\n transport input ssh\nabsent: ip http server\noptional: logging buffered 64000\n",
    template_aireos="config time ntp server 1 {{ntp}}\nabsent: config network telnet enable\n",
    template_variables={"ntp": "10.0.0.123"},
    max_clients_per_ap=50,
)

DNS = wa.DnsResults(
    forward={
        "HQ-B01-F01-AP001.corp.example": ["10.20.0.11"],
        "ap-lobby.corp.example": ["10.99.0.99"],
        "BR1-B01-F02-AP010.corp.example": ["10.40.0.10"],
        "br-wlc01.corp.example": ["10.1.1.1"],
    },
    reverse={
        "10.20.0.11": "hq-b01-f01-ap001.corp.example",
        "10.40.0.10": "something-else.corp.example",
        "10.0.0.5": "hq-wlc01.corp.example",
    },
)


def audit(standard: WirelessStandard = STANDARD, dns: wa.DnsResults | None = DNS, **kw) -> dict:
    facts = wa.gather(inputs(**kw), ignore_log_patterns=standard.log_ignore_patterns)
    return wa.run_audit(facts, standard, dns, now=NOW)


def findings(report: dict, check_id: str, controller: str | None = None) -> list[dict]:
    return [f for f in report["findings"] if f["check_id"] == check_id and (controller is None or f["controller_name"] == controller)]


def subjects(report: dict, check_id: str, controller: str | None = None) -> set:
    return {f["subject"] for f in findings(report, check_id, controller)}


def check_status(report: dict, controller: str, check_id: str) -> str:
    c = next(c for c in report["controllers"] if c["name"] == controller)
    return next(ch["status"] for ch in c["checks"] if ch["check_id"] == check_id)


# --- parsers ---------------------------------------------------------------------------------


def test_ap_summary_extra_reads_state_on_9800_and_clients_on_aireos():
    xe = wp.parse_ap_summary_extra(XE_AP_SUMMARY)
    assert xe["HQ-B01-F01-AP001"] == {"country": "US", "state": "Registered", "clients": None}
    assert xe["a4b2.3c4d.5e80"]["state"] == "Downloading"
    ae = wp.parse_ap_summary_extra(AIREOS_AP_SUMMARY)
    assert ae["BR1-B01-F02-AP010"] == {"country": "US", "state": None, "clients": 60}


def test_radio_summary_both_platforms():
    xe = {r["name"]: r for r in wp.parse_radio_summary(XE_5GHZ, "5")}
    assert xe["HQ-B01-F01-AP001"]["channel"] == 36
    assert xe["HQ-B01-F01-AP001"]["power_level"] == "2/8"
    assert xe["ap-lobby"]["oper"] == "Down"
    assert xe["a4b2.3c4d.5e80"]["admin"] == "Disabled"
    ae = {r["name"]: r for r in wp.parse_radio_summary(AIREOS_RADIOS_B, "2.4")}
    assert ae["BR1-B01-F02-AP010"]["channel"] == 11
    assert ae["BR1-B01-F02-AP011"]["admin"] == "Disabled"


def test_ap_ethernet_and_tags():
    eth = wp.parse_ap_ethernet(XE_ETHERNET, {"HQ-B01-F01-AP001", "ap-lobby"})
    assert eth["HQ-B01-F01-AP001"] == {"speed_mbps": 2500, "duplex": "Full", "up": True}
    assert eth["ap-lobby"]["duplex"] == "Half"
    tags = wp.parse_ap_tags(XE_TAGS)
    assert tags["ap-lobby"]["policy_tag"] == "default-policy-tag"
    assert tags["a4b2.3c4d.5e80"]["misconfigured"] is True


def test_client_health_detail_both_platforms():
    assert wp.parse_client_health(xe_client_detail(-60, 30)) == {"rssi": -60, "snr": 30, "policy_state": "RUN", "username": "alice", "vlan": None}
    ae = wp.parse_client_health(AIREOS_CLIENT_DETAIL)
    assert (ae["rssi"], ae["snr"], ae["policy_state"], ae["username"]) == (-55, 40, "RUN", "bob")


def test_controller_state_parsers():
    assert wp.parse_redundancy(XE_REDUNDANCY)["mode"] == "sso"
    assert wp.parse_redundancy(AIREOS_REDUNDANCY)["peer_state"] == "N/A"
    assert wp.parse_ntp_sync(XE_NTP) is True
    assert wp.parse_ntp_sync(AIREOS_TIME) is False
    assert wp.parse_cpu("CPU utilization for five seconds: 3%/0%; one minute: 2%; five minutes: 7%") == 7
    assert wp.parse_cpu("Current CPU(s) load: 2%") == 2
    assert wp.parse_software_version(XE_VERSION, "9800") == "17.09.05"
    assert wp.parse_software_version(AIREOS_SYSINFO, "aireos") == "8.10.185.0"


def test_syslog_grouping_and_traplog_shapes():
    groups = {g["key"]: g for g in wp.parse_syslog(XE_LOG)}
    assert groups["%CAPWAPAC_SMGR-4-AP_DISJOIN"]["count"] == 2
    assert groups["%CAPWAPAC_SMGR-4-AP_DISJOIN"]["notable"] is True
    assert groups["%SYS-2-MALLOCFAIL"]["severity"] == 2
    traps = wp.parse_traplog(AIREOS_TRAPLOG)
    assert traps[0]["key"] == "AP Disassociated" and traps[0]["count"] == 2 and traps[0]["notable"]


def test_wlan_security_flags():
    xe = {w["ssid"]: w for w in wp.parse_wlans(XE_WLANS, XE_RUN, "9800")}
    assert xe["GUEST"]["open"] and not xe["CORP"]["open"]
    assert xe["LEGACY"]["tkip"] and xe["LEGACY"]["wpa1"] and not xe["LEGACY"]["enabled"]
    ae = {w["ssid"]: w for w in wp.parse_wlans(AIREOS_WLANS, AIREOS_RUN_COMMANDS, "aireos")}
    assert ae["GUEST"]["open"] and ae["GUEST"]["webauth"]
    assert not ae["CORP"]["open"]


# --- golden template -------------------------------------------------------------------------


def test_template_presence_absence_nesting_and_variables():
    config = "hostname WLC1\nip http server\nline vty 0 4\n transport input ssh\nline vty 5 15\n transport input telnet\nntp server 10.0.0.1\n"
    template = "\n".join(
        [
            "! comment",
            "hostname {{device_name}}",
            "ntp server {{ntp}}",
            "absent: ip http server",
            "line vty {{*}}",
            " transport input ssh",
            "optional: banner motd {{*}}",
            "re: ^ntp server 10\\.0\\.0\\.\\d+$",
            "re: [unclosed",
        ]
    )
    result = ct.check_template(template, config, {"device_name": "WLC1", "ntp": "10.0.0.1"})
    by_kind = {(i.kind, i.rule, tuple(i.path)) for i in result.issues}
    assert ("forbidden", "absent: ip http server", ()) in by_kind
    assert ("missing", "transport input ssh", ("line vty 5 15",)) in by_kind
    assert ("missing", "optional: banner motd {{*}}", ()) in by_kind
    assert any(i.kind == "invalid" for i in result.issues)
    optional = next(i for i in result.issues if i.rule.startswith("optional:"))
    assert optional.severity == "warn"
    # hostname and both ntp lines are satisfied; vty 0 4 is fine.
    assert not any(i.rule in ("hostname {{device_name}}", "ntp server {{ntp}}") for i in result.issues)
    assert ct.check_template("hostname {{device_name}}", config, {"device_name": "OTHER"}).issues[0].kind == "missing"


def test_template_whitespace_and_case_insensitive():
    assert ct.check_template("NTP  Server 10.0.0.1", "ntp server   10.0.0.1").issues == []


# --- the audit -----------------------------------------------------------------------------


def test_never_collected_controller_fails_and_has_no_score():
    report = audit()
    new = next(c for c in report["controllers"] if c["name"] == "NEW-WLC")
    assert new["score"] is None
    assert findings(report, "data.snapshot", "NEW-WLC")[0]["severity"] == "fail"


def test_ap_naming_dns_addressing():
    report = audit()
    # Pattern + case + factory-default name.
    assert subjects(report, "ap.naming", "HQ-WLC01") == {"ap-lobby", "a4b2.3c4d.5e80"}
    assert any("factory default" in f["message"] for f in findings(report, "ap.naming"))
    assert subjects(report, "ap.naming", "BR-WLC01") == set()
    # Subnet and duplicate IP.
    assert subjects(report, "ap.subnet") == {"ap-lobby"}
    assert subjects(report, "ap.dup_ip") == {"HQ-B01-F01-AP001", "a4b2.3c4d.5e80"}
    # Forward DNS: mismatch for ap-lobby, missing record for the rest.
    fwd = {f["subject"]: f for f in findings(report, "ap.dns_forward")}
    assert "resolves to 10.99.0.99" in fwd["ap-lobby"]["message"]
    assert "HQ-B01-F01-AP001" not in fwd and "BR1-B01-F02-AP010" not in fwd
    assert "no DNS record" in fwd["BR1-B01-F02-AP011"]["message"]
    # Reverse DNS: wrong PTR is a fail, no PTR a warning.
    rev = {f["subject"]: f for f in findings(report, "ap.dns_reverse")}
    assert rev["BR1-B01-F02-AP010"]["severity"] == "fail"
    assert rev["BR1-B01-F02-AP011"]["severity"] == "warn"
    assert "HQ-B01-F01-AP001" not in rev
    row = next(a for a in report["access_points"] if a["name"] == "HQ-B01-F01-AP001")
    assert (row["name_ok"], row["subnet_ok"], row["dns_ok"], row["ptr_ok"]) == (True, True, True, True)
    assert row["uplink"] == "2500 Mbps Full"
    assert "5G ch36 Up" in row["radios"]


def test_ap_hardware_state_tags_and_load():
    report = audit()
    assert subjects(report, "ap.model") == {"ap-lobby"}
    assert subjects(report, "ap.software") == {"ap-lobby"}
    assert subjects(report, "ap.state") == {"a4b2.3c4d.5e80"}
    assert subjects(report, "ap.country") == {"ap-lobby"}
    uplink = {(f["subject"], f["severity"]) for f in findings(report, "ap.uplink")}
    assert uplink == {("ap-lobby", "fail"), ("ap-lobby", "warn")}  # half duplex, 100 Mbps
    tags = {f["subject"]: f["severity"] for f in findings(report, "ap.tags")}
    assert tags == {"ap-lobby": "warn", "a4b2.3c4d.5e80": "fail"}
    assert subjects(report, "ap.primary") == {"ap-lobby"}
    assert subjects(report, "ap.load") == {"BR1-B01-F02-AP010"}  # 60 clients on AireOS


def test_radios():
    report = audit()
    state = {(f["subject"], f["severity"]) for f in findings(report, "radio.state")}
    assert ("ap-lobby", "fail") in state  # enabled but down
    assert ("a4b2.3c4d.5e80", "warn") in state  # admin disabled
    assert ("BR1-B01-F02-AP011", "warn") in state
    assert subjects(report, "radio.channel24") == {"ap-lobby"}  # channel 3


def test_controller_checks():
    report = audit()
    assert findings(report, "ctl.software", "HQ-WLC01")[0]["actual"] == "17.09.05"
    assert findings(report, "ctl.software", "BR-WLC01") == []
    assert findings(report, "ctl.ha", "HQ-WLC01") == []
    assert findings(report, "ctl.ha", "BR-WLC01")[0]["severity"] == "fail"
    assert findings(report, "ctl.ntp") == []
    assert subjects(report, "ctl.syslog", "HQ-WLC01") == {"10.0.0.99"}
    assert subjects(report, "ctl.syslog", "BR-WLC01") == {"10.0.0.514"}
    assert subjects(report, "ctl.dns_servers") == {"10.0.0.53"}  # missing on AireOS only
    assert findings(report, "ctl.ntp_sync", "BR-WLC01")[0]["severity"] == "fail"
    assert findings(report, "ctl.ntp_sync", "HQ-WLC01") == []
    assert findings(report, "ctl.cpu", "HQ-WLC01")[0]["actual"] == "88%"
    assert findings(report, "ctl.dns", "HQ-WLC01") == []
    assert findings(report, "ctl.dns", "BR-WLC01") == []


def test_stale_snapshot_warns():
    report = audit(xe_age_hours=72)
    assert findings(report, "data.snapshot", "HQ-WLC01")[0]["severity"] == "warn"


def test_security_checks():
    report = audit()
    assert subjects(report, "sec.telnet", "HQ-WLC01") == {"line vty 5 15"}
    assert findings(report, "sec.telnet", "BR-WLC01")[0]["severity"] == "fail"
    assert findings(report, "sec.http", "HQ-WLC01")
    assert findings(report, "sec.http", "BR-WLC01") == []
    assert findings(report, "sec.snmp", "HQ-WLC01")
    assert findings(report, "sec.pwenc") == [] and findings(report, "sec.aaa") == [] and findings(report, "sec.enable") == []


def test_template_checks_per_platform():
    report = audit()
    xe = {f["message"] for f in findings(report, "cfg.template", "HQ-WLC01")}
    assert "Missing: transport input ssh (in 'line vty 5 15')" in xe
    assert "Must not be configured: ip http server" in xe
    assert any(m.startswith("Missing: optional: logging buffered") for m in xe)
    ae = findings(report, "cfg.template", "BR-WLC01")
    assert [f["message"] for f in ae] == ["Must not be configured: config network telnet enable"]


def test_wlans():
    report = audit()
    assert subjects(report, "wlan.required") == {"VOICE"}
    sec = {(f["controller_name"], f["subject"], f["severity"]) for f in findings(report, "wlan.security")}
    assert ("HQ-WLC01", "GUEST", "fail") in sec  # open, no webauth
    assert ("BR-WLC01", "GUEST", "warn") in sec  # open with web-auth
    assert ("HQ-WLC01", "LEGACY", "fail") in sec  # TKIP
    allowed = STANDARD.model_copy(update={"open_ssids_allowed": ["guest"]})
    assert not any(f["subject"] == "GUEST" and "open" in f["message"] for f in findings(audit(allowed), "wlan.security"))


def test_client_health():
    report = audit()
    clients = {c["mac"]: c for c in report["clients"]}
    poor = clients["1111.2222.3333"]
    assert (poor["rssi"], poor["snr"], poor["health"], poor["ip"], poor["username"]) == (-80, 12, "poor", "10.30.0.50", "alice")
    assert clients["4444.5555.6666"]["health"] == "unknown"
    assert "state IP Learn" in clients["4444.5555.6666"]["issues"]
    assert "no IP" in clients["4444.5555.6666"]["issues"]
    assert "self-assigned IP" in clients["7777.8888.9999"]["issues"]
    assert "legacy 802.11" in clients["7777.8888.9999"]["issues"]
    assert clients["a4bb.6d12.3456"]["health"] == "good"
    stuck = clients["a4bb.6d12.3457"]
    assert stuck["health"] == "fair" and stuck["state"] == "8021X_REQD"
    assert any("authentication" in f["message"] for f in findings(report, "client.state", "BR-WLC01"))
    hq = next(h for h in report["client_health"] if h["controller_name"] == "HQ-WLC01")
    assert (hq["total"], hq["poor"], hq["band_24"], hq["band_5"], hq["no_ip"], hq["legacy"]) == (3, 1, 2, 1, 2, 1)


def test_logs():
    report = audit()
    log = {f["subject"]: f["severity"] for f in findings(report, "log.syslog", "HQ-WLC01")}
    assert log == {"%SYS-2-MALLOCFAIL": "fail", "%DOT1X-3-SERVER_DEAD": "warn", "%CAPWAPAC_SMGR-4-AP_DISJOIN": "warn"}
    assert subjects(report, "log.traplog", "BR-WLC01") == {"AP Disassociated"}
    ignored = audit(STANDARD.model_copy(update={"log_ignore_patterns": ["MALLOCFAIL"]}))
    assert "%SYS-2-MALLOCFAIL" not in subjects(ignored, "log.syslog")


def test_without_dns_and_with_empty_standard():
    report = audit(WirelessStandard(), None)
    assert report["dns_checked"] is False
    assert check_status(report, "HQ-WLC01", "ap.dns_forward") == "skipped"
    # No pattern set: only the factory-default name is flagged.
    assert subjects(report, "ap.naming") == {"a4b2.3c4d.5e80"}
    assert check_status(report, "HQ-WLC01", "cfg.template") == "skipped"
    hq = next(c for c in report["controllers"] if c["name"] == "HQ-WLC01")
    assert 0 <= hq["score"] <= 100


def test_scores_and_summary():
    report = audit()
    hq = next(c for c in report["controllers"] if c["name"] == "HQ-WLC01")
    assert hq["platform"] == "9800" and hq["ap_count"] == 3 and hq["client_count"] == 3
    assert hq["score"] < 60
    assert report["summary"]["controllers"] == 3
    assert report["summary"]["fail"] > 0
    assert report["findings"][0]["severity"] == "fail"


def test_dns_targets():
    facts = wa.gather(inputs())
    names, ips = wa.dns_targets(facts, STANDARD)
    assert "HQ-B01-F01-AP001.corp.example" in names and "br-wlc01.corp.example" in names
    assert "10.40.0.11" in ips and "10.0.0.5" in ips


def test_ap_group_and_tag_parsers():
    groups = {g["name"]: g for g in wp.parse_ap_groups(AIREOS_APGROUPS)}
    assert groups["BR1-APGROUP01"]["wlans"] == ["1", "2"]
    assert groups["BR1-APGROUP01"]["rf_24"] == "RF24_BR" and groups["BR1-APGROUP01"]["rf_5"] is None
    assert [a["name"] for a in groups["default-group"]["aps"]] == ["BR1-B01-F02-AP011"]
    tags = wp.parse_tag_config(XE_TAG_CONFIG)
    assert tags["site"]["HQ01-APGROUP01-Flex"] == {"description": None, "wlans": {}, "ap_profile": None, "flex_profile": "HQ01_FLEXPROFILE01", "local_site": False, "rf": {}}
    assert tags["policy"]["pt-corp"]["wlans"] == {"CORP": "PP_WLAN_1"}
    assert tags["rf"]["rf-hq"]["rf"] == {"24ghz": "RF24_HQ", "5ghz": "RF5_HQ"}
    assert wp.parse_wncd_affinity(XE_AFFINITY)[0] == {"tag": "st-hq", "aps": 90, "instance": 0, "load": 0.18}
    assert wp.parse_wncd_cpu(XE_WNCD_CPU) == {0: 9, 1: 85, 2: 0}


def test_groups_and_tags_checks():
    report = audit()
    groups = {(g["controller_name"], g["kind"], g["name"]): g for g in report["groups"]}
    br = groups[("BR-WLC01", "AP group", "BR1-APGROUP01")]
    assert br["ap_count"] == 2 and br["wlans"] == "1, 2" and br["profiles"] == "2.4 GHz RF24_BR"
    assert groups[("HQ-WLC01", "Site tag", "HQ01-APGROUP01-Flex")]["profiles"] == "flex HQ01_FLEXPROFILE01, remote (no local-site)"
    assert groups[("HQ-WLC01", "Site tag", "st-hq")]["wncd"] == 0
    assert groups[("HQ-WLC01", "Policy tag", "pt-corp")]["wlans"] == "CORP → PP_WLAN_1"
    # Naming: st-hq and default tags/groups excluded from the pattern check.
    assert subjects(report, "grp.naming") == {"st-hq"}
    assert subjects(report, "grp.default", "BR-WLC01") == {"default-group"}
    assert {"default-site-tag", "default-policy-tag"} <= subjects(report, "grp.default", "HQ-WLC01")
    # 2802 and 9120 in one AP group.
    assert subjects(report, "grp.mixed", "BR-WLC01") == {"BR1-APGROUP01"}
    ap = next(a for a in report["access_points"] if a["name"] == "BR1-B01-F02-AP011")
    assert ap["ap_group"] == "default-group"
    # Defined but empty tag.
    assert "HQ01-APGROUP01-Flex" in subjects(report, "grp.unused", "HQ-WLC01")


def test_wncd_checks():
    report = audit()
    rows = {w["instance"]: w for w in report["wncd"] if w["controller_name"] == "HQ-WLC01"}
    assert rows[0]["aps"] == 90 and rows[0]["site_tags"] == ["st-hq"] and rows[1]["cpu_5min"] == 85
    bal = findings(report, "wncd.balance", "HQ-WLC01")
    assert any(f["subject"] == "wncd_0" and f["severity"] == "warn" for f in bal)
    assert subjects(report, "wncd.cpu", "HQ-WLC01") == {"wncd_1"}
    assert check_status(report, "BR-WLC01", "grp.naming") == "pass"


# --- API -------------------------------------------------------------------------------------

async def _register(client: AsyncClient, email: str) -> dict:
    resp = await client.post("/api/auth/register", json={"email": email, "password": "supersecret1", "org_name": "Wireless Org"})
    assert resp.status_code in (200, 201), resp.text
    token = resp.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.mark.asyncio
async def test_api_standard_audit_export_and_coverage(client: AsyncClient, unique_email, monkeypatch):
    headers = await _register(client, unique_email)

    resp = await client.get("/api/wireless-audit/standard", headers=headers)
    assert resp.status_code == 200
    assert resp.json()["is_default"] is True

    bad = STANDARD.model_dump()
    bad["ap_name_pattern"] = "([unclosed"
    resp = await client.put("/api/wireless-audit/standard", json=bad, headers=headers)
    assert resp.status_code == 400

    resp = await client.put("/api/wireless-audit/standard", json=STANDARD.model_dump(), headers=headers)
    assert resp.status_code == 200, resp.text
    assert resp.json()["is_default"] is False
    assert resp.json()["standard"]["ap_subnets"] == ["10.20.0.0/22", "10.40.0.0/22"]

    async def fake_inputs(db, org_id):
        return inputs()

    async def fake_resolve(names, ips, on_done=None):
        return DNS

    monkeypatch.setattr(wa_api, "_inputs", fake_inputs)
    monkeypatch.setattr(wa_api, "resolve", fake_resolve)

    resp = await client.get("/api/wireless-audit", headers=headers)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["dns_checked"] is True
    assert {c["name"] for c in body["controllers"]} == {"HQ-WLC01", "BR-WLC01", "NEW-WLC"}
    assert any(f["check_id"] == "ap.dns_forward" for f in body["findings"])

    progress = (await client.get("/api/wireless-audit/progress", headers=headers)).json()
    assert progress["running"] is False and progress["stage"] == "Done"
    lines = "\n".join(progress["lines"])
    assert "Parsing HQ-WLC01" in lines and "Looking up" in lines and "Checking BR-WLC01" in lines and "Checks done" in lines

    resp = await client.get("/api/wireless-audit/export.xlsx?dns=false", headers=headers)
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.content))
    assert wb.sheetnames == ["Summary", "Checks", "Findings", "Access points", "Clients", "Client health", "AP groups & tags", "WNCD", "Logs"]
    assert wb["Access points"].max_row == 6  # header + 5 APs

    resp = await client.get("/api/wireless-audit/collection-commands", headers=headers)
    assert resp.status_code == 200
    by_type = {c["device_type"]: c for c in resp.json()}
    # Everyday collections stay lean; the audit's collection adds the rest.
    xe = by_type["cisco_wlc_9800"]
    assert "show wireless client mac-address {client} detail" not in xe["base"]
    assert "show wireless client mac-address {client} detail" in xe["audit_added"] and "show ap tag summary" in xe["commands"]
    assert xe["commands"][: len(xe["base"])] == xe["base"]
    assert "show client detail {client}" in by_type["cisco_wlc"]["audit_added"]
    assert "show running-config" not in by_type["cisco_wlc"]["audit_added"]  # "show run-config" covers it


@pytest.mark.asyncio
async def test_resolve_bounds_lookups(monkeypatch):
    async def forward(name):
        return ["10.0.0.1"] if name == "ok" else []

    async def reverse(ip):
        return "host.example" if ip == "10.0.0.1" else None

    monkeypatch.setattr(wa_api, "_forward_lookup", forward)
    monkeypatch.setattr(wa_api, "_reverse_lookup", reverse)
    result = await wa_api.resolve({"ok", "missing"}, {"10.0.0.1", "10.0.0.2"})
    assert result.forward == {"ok": ["10.0.0.1"], "missing": []}
    assert result.reverse == {"10.0.0.1": "host.example", "10.0.0.2": None}


@pytest.mark.asyncio
async def test_real_lookups_use_their_own_pool():
    assert await wa_api._forward_lookup("localhost")
    assert await wa_api._forward_lookup("does-not-exist.invalid") == []



@pytest.mark.asyncio
async def test_dns_budget_marks_slow_lookups_unchecked(monkeypatch):
    import asyncio

    async def slow(name):
        await asyncio.sleep(10)
        return ["10.0.0.1"]

    async def fast(ip):
        return "host.example"

    monkeypatch.setattr(wa_api, "_forward_lookup", slow)
    monkeypatch.setattr(wa_api, "_reverse_lookup", fast)
    monkeypatch.setattr(wa_api, "DNS_BUDGET_SECONDS", 0.2)
    monkeypatch.setattr(wa_api, "DNS_TIMEOUT_SECONDS", 30)
    result = await wa_api.resolve({"ap1.corp"}, {"10.0.0.1"})
    assert result.unchecked == {"ap1.corp"} and result.reverse == {"10.0.0.1": "host.example"}
    # An unchecked AP is neither passed nor failed.
    dns = wa.DnsResults(forward={}, reverse={}, unchecked={"BR1-B01-F02-AP011.corp.example", "10.40.0.11"})
    report = wa.run_audit(wa.gather(inputs()), STANDARD, dns, now=NOW)
    row = next(a for a in report["access_points"] if a["name"] == "BR1-B01-F02-AP011")
    assert row["dns_ok"] is None and row["ptr_ok"] is None
